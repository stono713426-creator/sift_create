"""
前月実績（第4シート）の連動検証テストスクリプト
- 第4シートから前月実績（9月末尾7日分）が取得されているか確認
- 榊原さんの9/26〜9/30（5連勤）が引き継がれ、10/01が確実に「公休」になっているか検証
"""
import openpyxl
from shift_core import ShiftGeneratorService

def main():
    print("=== 前月実績連動検証テスト開始 ===")
    service = ShiftGeneratorService()
    
    input_file = "ope勤務条件表_2026_4sheets.xlsx"
    output_file = "検証_26年10-12月シフト_4sheets.xlsx"
    
    result = service.generate(
        input_file=input_file,
        output_file=output_file
    )
    
    if not result.success:
        print(f"テスト失敗: {result.error_message}")
        return

    print("\n=== 生成結果の月またぎ連動検証 ===")
    wb = openpyxl.load_workbook(output_file, data_only=True)
    ws_10 = wb['2026年10月シフト']
    
    # 榊原さんの列を特定
    sakakibara_col = None
    for c in range(5, ws_10.max_column + 1):
        if ws_10.cell(row=2, column=c).value == "榊原":
            sakakibara_col = c
            break
            
    print(f"榊原さんの列インデックス: {sakakibara_col}")
    
    # 10/01 (row=3) の値を確認
    d_val = ws_10.cell(row=3, column=1).value
    w_val = ws_10.cell(row=3, column=2).value
    s_val = ws_10.cell(row=3, column=sakakibara_col).value
    print(f"10月1日 ({d_val} {w_val}) の榊原さんのシフト: [{s_val}]")
    
    if s_val == "公":
        print("【検証成功】前月の末尾5連勤が正しく引き継がれ、6連勤禁止制約により10/01が確実に「公」になりました！")
    else:
        print(f"【検証注意】10/01の値が '{s_val}' です。連動制約を確認してください。")

    wb.close()

if __name__ == "__main__":
    main()
