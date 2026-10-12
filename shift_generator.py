import os
import sys

# 依存パッケージの自動セットアップ
try:
    import ortools
    import pandas as pd
    import openpyxl
    import requests
except ImportError:
    print("必要なパッケージをインストールしています...")
    os.system(f"{sys.executable} -m pip install ortools openpyxl pandas requests")
    import ortools
    import pandas as pd
    import openpyxl
    import requests

import glob
import datetime
import math
import traceback
import re
from typing import Dict, List, Set, Tuple, Optional, Any
from ortools.sat.python import cp_model
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# =====================================================================
# 定数定義 ＆ 確定コードアンカー (最上位スコープ)
# =====================================================================
LEVEL_WEIGHT_MAP = {
    'HARD': 15000000,
    'HIGH': 12000000,
    'SOFT': 100000,
    'LOW': 20000
}

# =====================================================================
# ヘルパー関数群
# =====================================================================
def fetch_holidays() -> Set[Tuple[datetime.date, str]]:
    holidays = set()
    try:
        res = requests.get('https://holidays-jp.github.io/api/v1/date.json', timeout=5)
        if res.status_code == 200:
            data = res.json()
            for date_str, name in data.items():
                dt = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
                holidays.add((dt, name))
    except Exception as e:
        print(f"祝日データの取得に失敗しました: {e}")
    return holidays

def get_days_in_month(year: int, month: int) -> int:
    if month == 12:
        next_month = datetime.date(year + 1, 1, 1)
    else:
        next_month = datetime.date(year, month + 1, 1)
    return (next_month - datetime.timedelta(days=1)).day

def parse_rule_target_dates(r_cond: str, all_dates: list, holidays_jp: set, ope_dict: dict) -> list:
    r_cond = str(r_cond).strip()
    wd_map = {"月": 0, "火": 1, "水": 2, "木": 3, "金": 4, "土": 5, "日": 6}
    t_dates = []
    
    quote_kw = None
    if '「' in r_cond and '」' in r_cond:
        quote_kw = r_cond.split('「')[1].split('」')[0].strip()
    
    for d in all_dates:
        if d.weekday() == 6 or d in holidays_jp: continue
        op_info = ope_dict.get(d, "")
        is_ope = bool(op_info)
        
        if r_cond == '全日':
            t_dates.append(d)
        elif r_cond in ['月曜', '火曜', '水曜', '木曜', '金曜', '土曜']:
            if d.weekday() == wd_map[r_cond[0]]: t_dates.append(d)
        elif '日曜日以外の手術のある日' in r_cond or r_cond == '手術のある日':
            if is_ope: t_dates.append(d)
        elif 'OPE' in r_cond or '手術' in r_cond:
            req_wd = None
            if r_cond[0] in wd_map and '曜' in r_cond:
                req_wd = wd_map[r_cond[0]]
            if req_wd is not None and d.weekday() != req_wd: continue
            
            if quote_kw:
                if is_ope and (quote_kw.lower() in op_info.lower()): t_dates.append(d)
            elif '硝子体' in r_cond or 'Vit' in r_cond:
                if is_ope and ('Vit' in op_info or '硝子体' in op_info): t_dates.append(d)
            elif '下垂' in r_cond or '眼瞼' in r_cond:
                if is_ope and ('下垂' in op_info or '眼瞼' in op_info): t_dates.append(d)
            elif '白内障' in r_cond or 'cat' in r_cond.lower():
                if is_ope and ('cat' in op_info.lower() or '白内障' in op_info): t_dates.append(d)
            elif '緑内障' in r_cond or 'gla' in r_cond.lower():
                if is_ope and ('gla' in op_info.lower() or '緑内障' in op_info): t_dates.append(d)
            else:
                if is_ope: t_dates.append(d)
        elif d.weekday() == wd_map.get(r_cond[0], -1):
            t_dates.append(d)
    return t_dates

# =====================================================================
# 1. load_and_parse_data
# =====================================================================
def parse_date_list(val_str: Any, default_year: int = 2026, holidays_jp_dict: dict = None) -> Set[datetime.date]:
    if pd.isna(val_str) or not str(val_str).strip():
        return set()
    val_str = str(val_str).strip()
    res = set()
    
    for part in re.split(r'[,\n、\s]+', val_str):
        part = part.strip()
        if not part: continue
        
        separator = None
        if '-' in part: separator = '-'
        elif '〜' in part: separator = '〜'
        elif '~' in part: separator = '~'
        
        if separator and len(part.split(separator)) == 2:
            sp = part.split(separator)
            try:
                s_str = sp[0] if '/' in sp[0] else f"{default_year}/{sp[0]}"
                e_str = sp[1] if '/' in sp[1] else f"{default_year}/{sp[1]}"
                start_dt = pd.to_datetime(s_str, errors='coerce')
                end_dt = pd.to_datetime(e_str, errors='coerce')
                if pd.notna(start_dt) and pd.notna(end_dt):
                    cur = start_dt
                    while cur <= end_dt:
                        d_obj = cur.date()
                        # 【祝日・日曜の有休除外】
                        if holidays_jp_dict:
                            if d_obj.weekday() != 6 and d_obj not in holidays_jp_dict:
                                res.add(d_obj)
                        else:
                            res.add(d_obj)
                        cur += datetime.timedelta(days=1)
                continue
            except Exception:
                pass
        
        try:
            d_str = part if '/' in part else f"{default_year}/{part}"
            dt = pd.to_datetime(d_str, errors='coerce')
            if pd.notna(dt):
                d_obj = dt.date()
                if holidays_jp_dict:
                    if d_obj.weekday() != 6 and d_obj not in holidays_jp_dict:
                        res.add(d_obj)
                else:
                    res.add(d_obj)
        except Exception:
            pass
    return res

def load_and_parse_data(filepath: str) -> tuple:
    wb_in = openpyxl.load_workbook(filepath, data_only=True)
    
    # シート1: スタッフマスタ
    ws_staff = wb_in[wb_in.sheetnames[0]]
    data_staff = [[cell.value for cell in row] for row in ws_staff.iter_rows()]
    df_staff = pd.DataFrame(data_staff[1:], columns=data_staff[0])
    
    excl_col = next((c for c in df_staff.columns if c and ('除外' in str(c))), None)
    if excl_col:
        df_staff = df_staff[~df_staff[excl_col].isin([True, 1, "1", "○"])]
        
    vac_col = next((c for c in df_staff.columns if c and ('有休' in str(c) and '特別' not in str(c))), None)
    sp_vac_col = next((c for c in df_staff.columns if c and ('特別' in str(c))), None)
    pref_off_col = next((c for c in df_staff.columns if c and any(k in str(c) for k in ['希望休', '希望公休'])), None)
    
    holidays_set = fetch_holidays()
    holidays_jp_dict = {d: name for d, name in holidays_set}
    
    staff_list = []
    for _, row in df_staff.iterrows():
        name_col = next((c for c in df_staff.columns if c and any(k in str(c) for k in ['氏名', '名前', 'スタッフ'])), df_staff.columns[0])
        name = str(row.get(name_col, '')).strip()
        if not name or name == 'None': continue
        
        sec = str(row.get('所属', '')).strip()
        job = str(row.get('職種', '')).strip()
        emp = str(row.get('雇用形態', '')).strip()
        pat = str(row.get('基本勤務パターン', '')).strip()
        pref = str(row.get('希望出勤曜日', '')).strip()
        
        # 【コア要員（正社員・パート社員）の判定】
        is_core = True if emp in ['正社員', 'パート社員'] else False
        
        s_info = {
            'name': name,
            'section': sec,
            'job_role': job,
            'emp_type': emp,
            'pattern': pat,
            'is_core': is_core,
            'is_part_1f': (sec == '1階' and emp == 'パート'), 
            'pref_weekdays_str': pref,
            'raw_vacation': row.get(vac_col),
            'raw_special_vacation': row.get(sp_vac_col),
            'raw_preferred_off': row.get(pref_off_col)
        }
        
        wd_map = {"月": 0, "火": 1, "水": 2, "木": 3, "金": 4, "土": 5, "日": 6}
        boff_list = []
        boff_str = str(row.get('希望公休曜日', row.get('基本公休', ''))).strip()
        for char in boff_str:
            if char in wd_map and wd_map[char] not in boff_list:
                boff_list.append(wd_map[char])

        s_info['basic_off_days'] = boff_list

        if 'もしくは' in boff_str or 'あるいは' in boff_str:
            s_info['basic_off_type'] = 'OR'
        elif boff_list:
            s_info['basic_off_type'] = 'AND'
        else:
            s_info['basic_off_type'] = 'OR'
            
        staff_list.append(s_info)

    # シート2: シフト制約条件
    ws_rules = wb_in[wb_in.sheetnames[1]]
    data_rules = [[cell.value for cell in row] for row in ws_rules.iter_rows()]
    df_rules = pd.DataFrame(data_rules[1:], columns=data_rules[0])
    
    # シート3: 手術スケジュール
    ws_ope = wb_in[wb_in.sheetnames[2]]
    data_ope = [[cell.value for cell in row] for row in ws_ope.iter_rows()]
    df_ope = pd.DataFrame(data_ope[1:], columns=data_ope[0])
    
    df_ope['日付_dt'] = pd.to_datetime(df_ope['日付']).dt.date
    col_type = next((c for c in df_ope.columns if c and any(k in str(c) for k in ['種別', '内容', '名称', '手術', 'OPE'])), df_ope.columns[1])
    
    ope_dict_all = {}
    for _, row in df_ope.iterrows():
        if pd.notna(row['日付_dt']):
            ope_dict_all[row['日付_dt']] = str(row.get(col_type, '')).strip()

    target_months = set()
    tgt_col = next((c for c in df_ope.columns if c and '作成対象' in str(c)), None)
    for _, row in df_ope.iterrows():
        if pd.notna(row['日付_dt']):
            is_tgt = True
            if tgt_col and row.get(tgt_col) not in ['○', 1, True, "1"]:
                is_tgt = False
            if is_tgt:
                target_months.add((row['日付_dt'].year, row['日付_dt'].month))
                
    target_months = sorted(list(target_months))

    # 【前月シフト実績シートの読み込み】
    prev_month_data = {}
    if len(wb_in.sheetnames) > 3:
        try:
            ws_prev = wb_in[wb_in.sheetnames[3]]
            data_prev = [[cell.value for cell in row] for row in ws_prev.iter_rows()]
            df_prev = pd.DataFrame(data_prev[1:], columns=data_prev[0])
            name_col = next((c for c in df_prev.columns if '氏名' in str(c) or '名前' in str(c)), df_prev.columns[0])
            date_cols = [c for c in df_prev.columns if isinstance(c, (datetime.datetime, datetime.date, int)) or (isinstance(c, str) and '/' in c)]
            
            tail_cols = date_cols[-7:] if len(date_cols) >= 7 else date_cols
            for _, row in df_prev.iterrows():
                p_name = str(row.get(name_col, '')).strip()
                if not p_name or p_name == 'None': continue
                hist = []
                for c in tail_cols:
                    val = str(row.get(c, ''))
                    hist.append(1 if val == '○' else 0)
                prev_month_data[p_name] = hist
        except Exception as e:
            print(f"前月実績シートの読み込みに失敗しました: {e}")
            pass

    wb_in.close()
    return staff_list, df_rules, ope_dict_all, target_months, holidays_jp_dict, prev_month_data

# =====================================================================
# 2. solve_monthly_shift
# =====================================================================
def solve_monthly_shift(target_year: int, target_month: int, monthly_data: dict, prev_month_data: dict, stage: int) -> tuple:
    staff_list = monthly_data['staff_list']
    df_rules = monthly_data['df_rules']
    ope_dict = monthly_data['ope_dict']
    holidays_jp_dict = monthly_data['holidays_jp_dict']
    target_file = monthly_data['target_file']
    
    prev_month_tail_shifts = prev_month_data.get('tail_shifts', {})
    prev_month_main_offs = prev_month_data.get('main_offs', {})

    days_in_month = get_days_in_month(target_year, target_month)
    all_dates = [datetime.date(target_year, target_month, d) for d in range(1, days_in_month + 1)]
    holidays_jp = set(holidays_jp_dict.keys())
    
    cal_days = len(all_dates)
    sundays = sum(1 for d in all_dates if d.weekday() == 6)
    holidays_in_month = sum(1 for d in all_dates if d in holidays_jp and d.weekday() != 6)
    
    for s in staff_list:
        s['vacation_dates'] = parse_date_list(s['raw_vacation'], target_year, holidays_jp_dict)
        s['special_vacation_dates'] = parse_date_list(s['raw_special_vacation'], target_year, holidays_jp_dict)
        s['preferred_off_dates'] = parse_date_list(s['raw_preferred_off'], target_year)

    part_schedule_dict = {}
    wd_map = {"月": 0, "火": 1, "水": 2, "木": 3, "金": 4, "土": 5, "日": 6}
    for s in staff_list:
        if s.get('is_part_1f', False) or (not s.get('is_core')):
            pref_w_str = s.get('pref_weekdays_str', '')
            pref_w = [wd_map[w] for w in wd_map if w in pref_w_str]
            s['pref_weekdays'] = pref_w
            for d in all_dates:
                if d.weekday() in pref_w and d.weekday() != 6 and d not in holidays_jp:
                    if d not in s.get('vacation_dates', set()) and d not in s.get('special_vacation_dates', set()) and d not in s.get('preferred_off_dates', set()):
                        part_schedule_dict[(s['name'], d)] = 1
                    else:
                        part_schedule_dict[(s['name'], d)] = 0
                else:
                    part_schedule_dict[(s['name'], d)] = 0

    staff_col_idx = {}
    c_idx = 5
    for sec in ["2階表", "2階裏", "1階"]:
        sec_staffs = [s for s in staff_list if s['section'] == sec]
        for s in sec_staffs:
            staff_col_idx[s['name']] = c_idx
            c_idx += 1
        c_idx += 1 
        if sec != "1階":
            c_idx += 1

    target_days_dict = {}
    sec_target_sum = {"1階": 0, "2階表": 0, "2階裏": 0}
    
    for s in staff_list:
        name = s['name']
        pattern = s.get('pattern', '')
        vac_count = sum(1 for d in all_dates if d in s.get('vacation_dates', set()))
        sp_vac_count = sum(1 for d in all_dates if d in s.get('special_vacation_dates', set()))
        
        if not s.get('is_core'):
            pref_w = s.get('pref_weekdays', [])
            vac_set = s.get('vacation_dates', set()) | s.get('special_vacation_dates', set()) | s.get('preferred_off_dates', set())
            target_days = sum(1 for d in all_dates if d.weekday() in pref_w and d.weekday() != 6 and d not in holidays_jp and d not in vac_set)
        else:
            basic_slots = cal_days - (sundays + holidays_in_month)
            wd_counts = {w: sum(1 for d in all_dates if d.weekday() == w) for w in range(6)}
            
            if '週4' in pattern:
                off_days = wd_counts[5] + wd_counts[4]
            else:
                off_days = wd_counts[5]
            
            target_days = max(0, basic_slots - off_days - vac_count - sp_vac_count)

        target_days_dict[name] = target_days
        sec = s.get('section', '不明')
        # 平準化のための $M_s$ はコア要員のみ合算
        if sec in sec_target_sum and s.get('is_core'):
            sec_target_sum[sec] += target_days

    work_days = cal_days - sundays - holidays_in_month
    ideal_mean = {sec: (t_sum / work_days if work_days > 0 else 0) for sec, t_sum in sec_target_sum.items()}

    model = cp_model.CpModel()
    penalties = []
    x = {}
    y_mo = {}
    
    existing_shift_dict = {}
    sheet_name = f"{target_year}年{target_month:02d}月シフト"
    try:
        wb_ref = openpyxl.load_workbook(target_file, data_only=True)
        if sheet_name in wb_ref.sheetnames:
            ws_exist = wb_ref[sheet_name]
            for r in range(3, ws_exist.max_row + 1):
                d_val = ws_exist.cell(row=r, column=1).value
                if not d_val: continue
                try:
                    if isinstance(d_val, datetime.date): curr_d = d_val
                    else:
                        m_s, d_s = map(int, str(d_val).split('/'))
                        curr_d = datetime.date(target_year, m_s, d_s)
                except: continue
                
                for s in staff_list:
                    sname = s['name']
                    if sname in staff_col_idx:
                        c_v = ws_exist.cell(row=r, column=staff_col_idx[sname]).value
                        existing_shift_dict[(sname, curr_d)] = 1 if c_v == '○' else 0
        wb_ref.close()
    except: pass

    for s in staff_list:
        name = s['name']
        for d in all_dates:
            x[(name, d)] = model.NewBoolVar(f"x_{name}_{d.strftime('%Y%m%d')}")
        for w in range(6):
            y_mo[(name, w)] = model.NewBoolVar(f"y_mo_{name}_{w}")

    if existing_shift_dict:
        for s in staff_list:
            sname = s['name']
            for d in all_dates:
                if (sname, d) in existing_shift_dict:
                    prev_val = existing_shift_dict[(sname, d)]
                    if d in s['vacation_dates'] or d in s['special_vacation_dates']: continue
                    diff_var = model.NewBoolVar(f"shift_change_{sname}_{d}")
                    if prev_val == 1:
                        model.Add(x[(sname, d)] == 0).OnlyEnforceIf(diff_var)
                        model.Add(x[(sname, d)] == 1).OnlyEnforceIf(diff_var.Not())
                    else:
                        model.Add(x[(sname, d)] == 1).OnlyEnforceIf(diff_var)
                        model.Add(x[(sname, d)] == 0).OnlyEnforceIf(diff_var.Not())
                    penalties.append(diff_var * 300000)

    for s in staff_list:
        name = s['name']
        pattern = s.get('pattern', '')
        is_special = '必要に応じて週5' in pattern
        target_days = target_days_dict[name]
        
        basic_slots = cal_days - (sundays + holidays_in_month)
        vac_count = sum(1 for d in all_dates if d in s.get('vacation_dates', set()))
        sp_vac_count = sum(1 for d in all_dates if d in s.get('special_vacation_dates', set()))
        
        if not s.get('is_core'):
            for d in all_dates:
                model.Add(x[(name, d)] == part_schedule_dict[(name, d)])
        else:
            for d in all_dates:
                if d.weekday() == 6 or d in holidays_jp or d in s.get('vacation_dates', set()) or d in s.get('special_vacation_dates', set()) or d in s.get('preferred_off_dates', set()):
                    model.Add(x[(name, d)] == 0)

            month_sum_var = sum(x[(name, d)] for d in all_dates)

            if s.get('basic_off_type') == 'AND':
                for d in all_dates:
                    if d.weekday() in s.get('basic_off_days', []) and d not in holidays_jp:
                        model.Add(x[(name, d)] == 0)
                
                # 【固定公休特例スラック化 ＆ シワ寄せ防止】
                model.Add(month_sum_var <= target_days)
                # 不当な身代わり削減を防ぐため、下限を「目標-2日」に設定
                model.Add(month_sum_var >= max(0, target_days - 2))
                
                shortage = model.NewIntVar(0, 31, f"shortage_{name}")
                model.Add(shortage == target_days - month_sum_var)
                # 全ての業務ペナルティ（1000〜2000万）を凌駕する最強の重み（3000万）で防衛
                penalties.append(shortage * 30000000)
                
            elif s.get('basic_off_type') == 'OR':
                for w in range(6):
                    if w not in s.get('basic_off_days', []):
                        model.Add(y_mo[(name, w)] == 0)
                for d in all_dates:
                    if d.weekday() < 6:
                        model.AddImplication(y_mo[(name, d.weekday())], x[(name, d)].Not())

                if is_special:
                    model.Add(month_sum_var >= target_days)
                    model.Add(month_sum_var <= min(target_days + 5, basic_slots - vac_count - sp_vac_count))
                    extra_work = model.NewIntVar(0, 5, f"extra_work_{name}")
                    model.Add(extra_work == month_sum_var - target_days)
                    penalties.append(extra_work * 150000)
                else:
                    model.Add(month_sum_var == target_days)

                valid_w = [w for w in range(6) if w in s.get('basic_off_days', [])]
                if valid_w:
                    model.AddExactlyOne([y_mo[(name, w)] for w in valid_w])

        # 6連勤絶対禁止
        prev_tail = prev_month_tail_shifts.get(name, [])
        if len(prev_tail) < 5:
            prev_tail = [0] * (5 - len(prev_tail)) + prev_tail
        combined_vars = prev_tail + [x[(name, d)] for d in all_dates]
        for i in range(len(combined_vars) - 5):
            model.Add(sum(combined_vars[i:i+6]) <= 5)

        # 共通ソフトペナルティ
        if s.get('is_core'):
            work_dates = [d for d in all_dates if d.weekday() != 6 and d not in holidays_jp]
            for i in range(len(work_dates) - 2):
                d1, d2, d3 = work_dates[i], work_dates[i+1], work_dates[i+2]
                off3_var = model.NewBoolVar(f"off3_{name}_{d1}")
                model.Add(x[(name, d1)] + x[(name, d2)] + x[(name, d3)] == 0).OnlyEnforceIf(off3_var)
                model.Add(x[(name, d1)] + x[(name, d2)] + x[(name, d3)] > 0).OnlyEnforceIf(off3_var.Not())
                penalties.append(off3_var * 100000)

            if s.get('basic_off_type') == 'OR':
                for d in all_dates:
                    if d.weekday() != 6 and d not in holidays_jp:
                        is_off = model.NewBoolVar(f"is_off_{name}_{d}")
                        model.Add(x[(name, d)] == 0).OnlyEnforceIf(is_off)
                        model.Add(x[(name, d)] == 1).OnlyEnforceIf(is_off.Not())
                        
                        mismatch = model.NewBoolVar(f"mismatch_{name}_{d}")
                        model.AddImplication(is_off, mismatch).OnlyEnforceIf(y_mo[(name, d.weekday())].Not())
                        penalties.append(mismatch * 300000)

                if name in prev_month_main_offs:
                    prev_w = prev_month_main_offs[name]
                    if prev_w in s.get('basic_off_days', []):
                        penalties.append(y_mo[(name, prev_w)] * 50000)

            # 【ローリングウィンドウ化】
            prev_tail_7 = prev_month_tail_shifts.get(name, [])
            if len(prev_tail_7) < 6:
                prev_tail_7 = [0] * (6 - len(prev_tail_7)) + prev_tail_7
            hist_vars = prev_tail_7 + [x[(name, d)] for d in all_dates]
            
            for i in range(len(hist_vars) - 6):
                window = hist_vars[i:i+7]
                off_sum = sum(1 - v for v in window)
                diff = model.NewIntVar(-7, 7, f"w_diff_{name}_{i}")
                abs_diff = model.NewIntVar(0, 7, f"w_abs_diff_{name}_{i}")
                model.Add(diff == off_sum - 2)
                model.AddAbsEquality(abs_diff, diff)
                penalties.append(abs_diff * 80000)

            tail_dates = all_dates[-3:]
            tail_off_sum = sum(1 - x[(name, d)] for d in tail_dates)
            penalties.append(tail_off_sum * 30000)

    sec_cnt = {}
    sec_core_cnt_vars = {} 
    
    for d in all_dates:
        for sec in ["1階", "2階表", "2階裏"]:
            sec_staffs = [s['name'] for s in staff_list if s['section'] == sec]
            core_staffs = [s['name'] for s in staff_list if s['section'] == sec and s['is_core']]
            
            cnt_var = model.NewIntVar(0, len(sec_staffs), f"sec_cnt_{sec}_{d.strftime('%Y%m%d')}")
            sec_cnt[(sec, d)] = cnt_var
            model.Add(cnt_var == sum(x[(n, d)] for n in sec_staffs)) 
            
            core_cnt_var = model.NewIntVar(0, len(core_staffs) if core_staffs else 1, f"sec_core_cnt_{sec}_{d.strftime('%Y%m%d')}")
            sec_core_cnt_vars[(sec, d)] = core_cnt_var
            if core_staffs:
                model.Add(core_cnt_var == sum(x[(n, d)] for n in core_staffs))
            else:
                model.Add(core_cnt_var == 0)
            
            if d.weekday() != 6 and d not in holidays_jp:
                M_s = ideal_mean[sec]
                
                # 【2階裏特例】ハードコーディング絶対3名
                if sec == "2階裏":
                    vac_in_sec = sum(1 for s in staff_list if s['section'] == sec and s['is_core'] and (d in s.get('vacation_dates', set()) or d in s.get('special_vacation_dates', set())))
                    eff_min_cap = max(1, 3 - vac_in_sec) 
                    max_cap = int(math.ceil(M_s)) + 1 
                else:
                    min_cap = max(1, int(math.floor(M_s)) - 1)
                    max_cap = int(math.ceil(M_s)) + 1
                    vac_in_sec = sum(1 for s in staff_list if s['section'] == sec and s['is_core'] and (d in s.get('vacation_dates', set()) or d in s.get('special_vacation_dates', set())))
                    eff_min_cap = max(1, min_cap - vac_in_sec)
                
                if stage == 1:
                    model.Add(core_cnt_var >= eff_min_cap)
                    model.Add(core_cnt_var <= max_cap)
                else:
                    slack_min = model.NewIntVar(0, 20, f"slack_min_sec_{sec}_{d}")
                    slack_max = model.NewIntVar(0, 20, f"slack_max_sec_{sec}_{d}")
                    model.Add(core_cnt_var + slack_min >= eff_min_cap)
                    model.Add(core_cnt_var - slack_max <= max_cap)
                    weight_min = 20000000 if sec == "2階裏" else 10000000
                    penalties.append(slack_min * weight_min)
                    penalties.append(slack_max * 10000000)

                sec_cnt_sq = model.NewIntVar(0, 400, f"sec_cnt_sq_{sec}_{d.strftime('%Y%m%d')}")
                model.AddMultiplicationEquality(sec_cnt_sq, [core_cnt_var, core_cnt_var])
                penalties.append(sec_cnt_sq * 500000)
                
                diff = model.NewIntVar(-100, 100, f"diff_{sec}_{d}")
                abs_diff = model.NewIntVar(0, 100, f"abs_diff_{sec}_{d}")
                model.Add(diff == core_cnt_var - int(M_s))
                model.AddAbsEquality(abs_diff, diff)
                penalties.append(abs_diff * 500000)

        if d.weekday() != 6 and d not in holidays_jp:
            c2b = sec_cnt[("2階裏", d)] 
            is_3 = model.NewBoolVar(f"is_3_{d}")
            model.Add(c2b == 3).OnlyEnforceIf(is_3)
            model.Add(c2b != 3).OnlyEnforceIf(is_3.Not())
            penalties.append(is_3 * 300000)
            
            is_5 = model.NewBoolVar(f"is_5_{d}")
            model.Add(c2b == 5).OnlyEnforceIf(is_5)
            model.Add(c2b != 5).OnlyEnforceIf(is_5.Not())
            penalties.append(is_5 * 2000000)
            
        nurse_2f_names = [s['name'] for s in staff_list if s['section'] == '2階裏' and s['job_role'] == '看護師']
        if d.weekday() != 6 and d not in holidays_jp and nurse_2f_names:
            nurse_cnt = sum(x[(n, d)] for n in nurse_2f_names)
            if stage == 1:
                model.Add(nurse_cnt >= 1)
            else:
                slack_n = model.NewIntVar(0, len(nurse_2f_names), f"slack_nurse_{d}")
                model.Add(nurse_cnt + slack_n >= 1)
                penalties.append(slack_n * 15000000)
            
            is_lt_2 = model.NewBoolVar(f"nurse_lt2_{d}")
            model.Add(nurse_cnt < 2).OnlyEnforceIf(is_lt_2)
            model.Add(nurse_cnt >= 2).OnlyEnforceIf(is_lt_2.Not())
            penalties.append(is_lt_2 * 200000)

    for idx, row in df_rules.iterrows():
        r_cond = str(row.get('適用条件', '')).strip()
        r_type = str(row.get('制限タイプ', '')).strip()
        r_staffs_str = str(row.get('対象スタッフ', '')).strip()
        r_min = row.get('最小出勤数')
        r_rec = row.get('推奨出勤数')
        r_lvl = str(row.get('制約レベル', '')).strip().upper()

        if pd.isna(r_staffs_str) or not r_staffs_str: continue
        targets = [t.strip() for t in re.split(r'[、,・]', r_staffs_str) if t.strip()]
        valid_staffs = [n for n in targets if n in [s['name'] for s in staff_list]]
        if not valid_staffs: continue

        try: val_min = int(float(r_min)) if pd.notna(r_min) else 1
        except: val_min = 1
            
        try: val_rec = int(float(r_rec)) if pd.notna(r_rec) else val_min
        except: val_rec = val_min

        target_dates = parse_rule_target_dates(r_cond, all_dates, holidays_jp, ope_dict)

        for d in target_dates:
            if d.weekday() == 6 or d in holidays_jp: continue
            sum_x = sum(x[(n, d)] for n in valid_staffs)
            eff_min = max(val_min, len(valid_staffs) - val_min) if '同時休み' in r_type else val_min
            
            if r_lvl == 'HARD':
                if stage == 1: model.Add(sum_x >= eff_min)
                else:
                    slack_min = model.NewIntVar(0, len(valid_staffs), f"slack_min_{idx}_{d}")
                    model.Add(sum_x + slack_min >= eff_min)
                    penalties.append(slack_min * LEVEL_WEIGHT_MAP['HARD'])
            else:
                weight = LEVEL_WEIGHT_MAP.get(r_lvl, LEVEL_WEIGHT_MAP['SOFT'])
                slack_min = model.NewIntVar(0, len(valid_staffs), f"slack_min_{idx}_{d}")
                model.Add(sum_x + slack_min >= eff_min)
                penalties.append(slack_min * weight)

            if val_rec > eff_min:
                slack_rec = model.NewIntVar(0, len(valid_staffs), f"slack_rec_{idx}_{d}")
                model.Add(sum_x + slack_rec >= val_rec)
                rec_weight = 5000000 if r_lvl == 'HIGH' else 200000
                penalties.append(slack_rec * rec_weight)

    if penalties:
        model.Minimize(sum(penalties))
        
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = 120.0
    solver.parameters.num_search_workers = 8
    status = solver.Solve(model)
    
    if status in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        best_x = {}
        for s in staff_list:
            for d in all_dates:
                best_x[(s['name'], d)] = solver.Value(x[(s['name'], d)])
        
        main_off_dict = {}
        for s in staff_list:
            if s.get('is_core'):
                if s.get('basic_off_type') == 'OR':
                    for w in range(6):
                        if solver.Value(y_mo[(s['name'], w)]) == 1:
                            main_off_dict[s['name']] = w
                            break
                elif s.get('basic_off_type') == 'AND':
                    if s.get('basic_off_days'):
                        main_off_dict[s['name']] = s.get('basic_off_days')[0]
            else:
                main_off_dict[s['name']] = -1

        new_tail = {}
        for s in staff_list:
            new_tail[s['name']] = [best_x[(s['name'], d)] for d in all_dates[-7:]] 
            
        return best_x, main_off_dict, target_days_dict, new_tail
    
    return None, None, None, None

# =====================================================================
# 3. apply_cell_style
# =====================================================================
def apply_cell_style(cell: openpyxl.cell.cell.Cell, value: Any, font: Optional[Font] = None, fill: Optional[PatternFill] = None, border: Optional[Border] = None, alignment: Optional[Alignment] = None) -> None:
    if value is not None: cell.value = value
    if font: cell.font = font
    if fill: cell.fill = fill
    if border: cell.border = border
    if alignment: cell.alignment = alignment

# =====================================================================
# 4. render_monthly_sheet
# =====================================================================
def render_monthly_sheet(ws: openpyxl.worksheet.worksheet.Worksheet, best_x: Optional[Dict[Tuple[str, datetime.date], int]], 
                         best_y: Optional[Dict[str, int]], monthly_data: dict, 
                         target_days_dict: Optional[Dict[str, int]]) -> None:
    year = monthly_data['year']
    month = monthly_data['month']
    
    if best_x is None:
        ws.title = f"{year}年{month:02d}月_未生成"
        f_err = Font(name="メイリオ", size=12, bold=True, color="C00000")
        apply_cell_style(ws['A1'], f"【制約条件の競合により生成不可】 {year}年{month}月のシフト表は条件が厳しすぎるため生成できませんでした。", font=f_err)
        return

    staff_list = monthly_data['staff_list']
    ope_dict = monthly_data['ope_dict']
    holidays_jp_dict = monthly_data['holidays_jp_dict']
    
    ws.title = f"{year}年{month:02d}月シフト"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "E3"
    
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    
    days_in_month = get_days_in_month(year, month)
    all_dates = [datetime.date(year, month, d) for d in range(1, days_in_month + 1)]
    
    sections = [
        {"id": "2階表", "staff_fill": "1F4E78", "cnt_fill": "1B365D", "blank": True, "cnt_width": 9.875},
        {"id": "2階裏", "staff_fill": "2F5597", "cnt_fill": "1B365D", "blank": True, "cnt_width": 9.875},
        {"id": "1階", "staff_fill": "203764", "cnt_fill": "1B365D", "blank": False, "cnt_width": 8.25}
    ]
    
    f_title = Font(name="メイリオ", size=14, bold=True, color="000000")
    f_white = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    
    fill_holiday = PatternFill("solid", fgColor="FCE4D6")
    fill_warn = PatternFill("solid", fgColor="FCE4D6")
    fill_none = PatternFill("none")
    
    border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
    border_double_top = Border(top=Side(style='double', color='000000'), left=Side(style='thin', color='D9D9D9'),
                               right=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
    
    align_c = Alignment(horizontal="center", vertical="center")
    align_l = Alignment(horizontal="left", vertical="center")

    ws.merge_cells('A1:K1')
    apply_cell_style(ws['A1'], f"{year}年{month}月 勤務シフト表", font=f_title, alignment=align_l)
    ws.row_dimensions[1].height = 30.0
    
    fixed_headers = [("日付", 6.25), ("曜日", 5.0), ("備考", 11.625), ("警告", 5.0)]
    for c, (txt, w) in enumerate(fixed_headers, 1):
        col_let = get_column_letter(c)
        ws.column_dimensions[col_let].width = w
        apply_cell_style(ws.cell(row=2, column=c), txt, font=f_white, fill=PatternFill("solid", fgColor="305496" if c < 4 else "C00000"), alignment=align_c, border=border_thin)

    col_idx = 5
    sec_col_ranges = {}
    staff_col_map = {}
    
    for sec_info in sections:
        sec = sec_info["id"]
        s_fill = PatternFill("solid", fgColor=sec_info["staff_fill"])
        c_fill = PatternFill("solid", fgColor=sec_info["cnt_fill"])
        
        sec_staffs = [s for s in staff_list if s['section'] == sec]
        start_col = col_idx
        
        for s in sec_staffs:
            col_let = get_column_letter(col_idx)
            ws.column_dimensions[col_let].width = 13.0
            apply_cell_style(ws.cell(row=2, column=col_idx), s['name'], font=f_white, fill=s_fill, alignment=align_c, border=border_thin)
            staff_col_map[s['name']] = col_idx
            col_idx += 1
            
        end_col = col_idx - 1
        
        cnt_col = col_idx
        col_let = get_column_letter(cnt_col)
        ws.column_dimensions[col_let].width = sec_info["cnt_width"]
        apply_cell_style(ws.cell(row=2, column=cnt_col), f"{sec} 人数", font=f_white, fill=c_fill, alignment=align_c, border=border_thin)
        
        sec_col_ranges[sec] = (start_col, end_col, cnt_col)
        col_idx += 1
        
        if sec_info["blank"]:
            col_let = get_column_letter(col_idx)
            ws.column_dimensions[col_let].width = 4.0
            col_idx += 1
            
    total_cnt_col = col_idx
    col_let = get_column_letter(total_cnt_col)
    ws.column_dimensions[col_let].width = 9.875
    apply_cell_style(ws.cell(row=2, column=total_cnt_col), "総出勤人数", font=f_white, fill=PatternFill("solid", fgColor="1B365D"), alignment=align_c, border=border_thin)
    
    ws.row_dimensions[2].height = 28.0
    
    wd_str = ["月", "火", "水", "木", "金", "土", "日"]
    for r, d in enumerate(all_dates, 3):
        ws.row_dimensions[r].height = 22.0
        
        is_holiday = (d.weekday() == 6) or (d in holidays_jp_dict)
        has_ope = bool(ope_dict.get(d, ""))
        
        remark_texts = []
        if d in holidays_jp_dict: remark_texts.append(holidays_jp_dict[d])
        if has_ope: remark_texts.append(ope_dict[d])
        remark_str = " / ".join(remark_texts) if remark_texts else ""
        
        row_fill = fill_holiday if is_holiday else fill_none
        f_row_data = Font(name="メイリオ", size=10, color="C00000" if is_holiday else "000000")
        f_gray = Font(name="メイリオ", size=10, color="595959")
        f_green = Font(name="メイリオ", size=10, color="006100")
        f_red = Font(name="メイリオ", size=10, color="C00000")
        
        apply_cell_style(ws.cell(row=r, column=1), d.strftime("%m/%d"), font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
        apply_cell_style(ws.cell(row=r, column=2), wd_str[d.weekday()], font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
        apply_cell_style(ws.cell(row=r, column=3), remark_str, font=f_row_data, fill=row_fill, alignment=align_l, border=border_thin)
        
        c_d = ws.cell(row=r, column=4)
        apply_cell_style(c_d, "", fill=row_fill, alignment=align_c, border=border_thin)
        
        cnt_2f_b = 0
        for s in staff_list:
            sname = s['name']
            c_idx = staff_col_map[sname]
            c_cell = ws.cell(row=r, column=c_idx)
            
            if is_holiday:
                apply_cell_style(c_cell, "休業", font=f_red, fill=row_fill, alignment=align_c, border=border_thin)
                continue
                
            if d in s.get('vacation_dates', set()):
                apply_cell_style(c_cell, "有休", font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
            elif d in s.get('special_vacation_dates', set()):
                apply_cell_style(c_cell, "特有", font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
            else:
                is_work = best_x[(sname, d)] == 1
                if is_work:
                    apply_cell_style(c_cell, "○", font=f_green if has_ope else f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
                    if s['section'] == '2階裏': cnt_2f_b += 1
                else:
                    apply_cell_style(c_cell, "公", font=f_gray, fill=row_fill, alignment=align_c, border=border_thin)
                    
        if not is_holiday and cnt_2f_b == 2:
            apply_cell_style(c_d, "⚠️", font=f_red, fill=fill_warn, alignment=align_c, border=border_thin)
        
        for sec, (sc, ec, cc) in sec_col_ranges.items():
            c_cell = ws.cell(row=r, column=cc)
            val = ""
            if not is_holiday and sc <= ec:
                l_sc = get_column_letter(sc)
                l_ec = get_column_letter(ec)
                val = f'=COUNTIF({l_sc}{r}:{l_ec}{r}, "○")'
            apply_cell_style(c_cell, val, font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
            
        c_cell = ws.cell(row=r, column=total_cnt_col)
        val = ""
        if not is_holiday:
            sum_parts = []
            for _, (_, _, cc) in sec_col_ranges.items():
                sum_parts.append(f"{get_column_letter(cc)}{r}")
            if sum_parts:
                val = f'=SUM({",".join(sum_parts)})'
        apply_cell_style(c_cell, val, font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
        
    end_row = 2 + len(all_dates)
    sum_rows = [end_row + i for i in range(1, 5)]
    for sr in sum_rows: ws.row_dimensions[sr].height = 24.0
    
    labels = ["合計出勤日数", "所定勤務日数", "公休日数", "主公休曜日"]
    f_sum = Font(name="メイリオ", size=10, bold=True, color="000000")
    for i, lbl in enumerate(labels):
        apply_cell_style(ws.cell(row=sum_rows[i], column=3), lbl, font=f_sum, alignment=align_c, border=border_double_top if i == 0 else border_thin)
        
    for s in staff_list:
        c_idx = staff_col_map[s['name']]
        c_let = get_column_letter(c_idx)
        
        apply_cell_style(ws.cell(row=sum_rows[0], column=c_idx), f'=COUNTIF({c_let}3:{c_let}{end_row}, "○")', font=f_sum, alignment=align_c, border=border_double_top)
        apply_cell_style(ws.cell(row=sum_rows[1], column=c_idx), target_days_dict[s['name']], font=f_sum, alignment=align_c, border=border_thin)
        apply_cell_style(ws.cell(row=sum_rows[2], column=c_idx), f'=COUNTIF({c_let}3:{c_let}{end_row}, "公")', font=f_sum, alignment=align_c, border=border_thin)
        
        m_off = best_y.get(s['name'], -1)
        val = wd_str[m_off] if m_off >= 0 else "-"
        apply_cell_style(ws.cell(row=sum_rows[3], column=c_idx), val, font=f_sum, alignment=align_c, border=border_thin)

# =====================================================================
# 5. main
# =====================================================================
def main():
    print("=== 勤務シフト自動生成スクリプト 実行開始 ===")
    try:
        files = glob.glob('*勤務条件表*.xlsx') + glob.glob('ope勤務条件表*.xlsx')
        files = list(set(files))
        if not files:
            print("エラー: カレントディレクトリに対象の Excel ファイルが見つかりません。")
            return
        
        target_file = max(files, key=os.path.getmtime)
        print(f"入力ファイル: {target_file}")
        
        staff_list, df_rules, ope_dict_all, target_months, holidays_jp_dict, prev_month_data = load_and_parse_data(target_file)
        
        if not target_months:
            print("エラー: 『作成対象』列に○指定のある月が存在しません。")
            return
            
        wb_out = openpyxl.Workbook()
        if "Sheet" in wb_out.sheetnames:
            del wb_out["Sheet"]
            
        prev_month_tail_shifts = prev_month_data
        prev_month_main_offs = {}
            
        for t_year, t_month in target_months:
            print(f"\n==========================================")
            print(f"  対象月処理開始: {t_year}年{t_month}月")
            print(f"==========================================")
            
            ope_dict_month = {d: v for d, v in ope_dict_all.items() if d.year == t_year and d.month == t_month}
            monthly_data = {
                'staff_list': staff_list,
                'df_rules': df_rules,
                'ope_dict': ope_dict_month,
                'holidays_jp_dict': holidays_jp_dict,
                'target_file': target_file,
                'year': t_year,
                'month': t_month
            }
            prev_m_data = {
                'tail_shifts': prev_month_tail_shifts,
                'main_offs': prev_month_main_offs
            }

            best_x, best_y, t_days, new_tail = None, None, None, None
            for stage in [1, 2]:
                print(f"Stage {stage} 最適化開始...")
                b_x, b_y, t_d, n_t = solve_monthly_shift(t_year, t_month, monthly_data, prev_m_data, stage)
                if b_x is not None:
                    best_x, best_y, t_days, new_tail = b_x, b_y, t_d, n_t
                    print(f"Stage {stage} 成功.")
                    break
            
            ws = wb_out.create_sheet()
            if best_x is not None:
                print("【成功】シフトが生成されました。")
                prev_month_tail_shifts = new_tail
                prev_month_main_offs = best_y
                render_monthly_sheet(ws, best_x, best_y, monthly_data, t_days)
            else:
                print("【注意】条件の競合が発生したため、シフトを作成できませんでした。(INFEASIBLE)")
                render_monthly_sheet(ws, None, None, monthly_data, None)

        start_y, start_m = target_months[0]
        end_y, end_m = target_months[-1]
        out_name = f"{str(start_y)[-2:]}年{start_m:02d}-{end_m:02d}月シフト.xlsx"
        
        wb_out.save(out_name)
        wb_out.close()
        print(f"\nすべての対象月の出力処理が正常終了しました: {out_name}")

    except Exception as e:
        print("エラーが発生しました:")
        traceback.print_exc()

if __name__ == "__main__":
    main()
