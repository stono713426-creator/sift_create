"""
データパースモジュール
スタッフマスタ、シフト制約条件、手術スケジュール、前月実績（縦型・横型両対応）の読み込み。
"""
import re
import numbers
import unicodedata
import datetime
from typing import Dict, List, Set, Tuple, Optional, Any
import openpyxl
import pandas as pd
from .holidays import fetch_holidays
from .config import ShiftConfig

def parse_settings_sheet(wb_in: openpyxl.Workbook) -> ShiftConfig:
    """『設定』または『システム設定』シートから動的パラメータを読み込み"""
    config = ShiftConfig()
    
    # シート検索
    target_sheet_name = next(
        (s for s in wb_in.sheetnames if any(k in s for k in ['設定', 'システム設定', 'パラメータ', 'config'])),
        None
    )
    if not target_sheet_name:
        return config

    try:
        ws_cfg = wb_in[target_sheet_name]
        data = [[cell.value for cell in row] for row in ws_cfg.iter_rows()]
        if len(data) < 2:
            return config

        # ヘッダー行検索
        header_row_idx = 0
        for r_idx in range(min(4, len(data))):
            row_str = [str(c) for c in data[r_idx] if c is not None]
            if any('項目' in s or '設定' in s for s in row_str) and any('値' in s for s in row_str):
                header_row_idx = r_idx
                break

        headers = [str(c).strip() if c is not None else "" for c in data[header_row_idx]]
        key_col = next((i for i, h in enumerate(headers) if any(k in h for k in ['項目', 'キー', '名前'])), 0)
        val_col = next((i for i, h in enumerate(headers) if any(k in h for k in ['値', '設定値'])), 1)

        raw_settings = {}
        for r in range(header_row_idx + 1, len(data)):
            row = data[r]
            if len(row) > max(key_col, val_col):
                k = row[key_col]
                v = row[val_col]
                if k is not None and v is not None and str(k).strip():
                    raw_settings[str(k).strip()] = v

        # 1. 2階裏 基本人数
        for k in ['2階裏_基本人数', '2階裏基本人数', '2階裏_基準人数']:
            if k in raw_settings:
                try: config.sec_2f_back_base_cap = int(float(raw_settings[k]))
                except: pass

        # 2. 2階裏 看護師最低人数
        for k in ['2階裏_看護師最低人数', '看護師最低人数', '2階裏看護師最低']:
            if k in raw_settings:
                try: config.nurse_min_count = int(float(raw_settings[k]))
                except: pass

        # 3. 2階裏 看護師推奨人数
        for k in ['2階裏_看護師推奨人数', '看護師推奨人数', '2階裏看護師推奨']:
            if k in raw_settings:
                try: config.nurse_rec_count = int(float(raw_settings[k]))
                except: pass

        # 4. 連勤上限日数
        for k in ['連勤上限日数', '連勤上限', '最大連勤数']:
            if k in raw_settings:
                try: config.max_consecutive_work_days = int(float(raw_settings[k]))
                except: pass

        # 5. 固定公休 削減許容日数
        for k in ['固定公休_削減許容日数', '固定公休削減許容', '固定公休不足許容']:
            if k in raw_settings:
                try: config.fixed_off_shortage_allowance = int(float(raw_settings[k]))
                except: pass

        # 6. 3連休抑制レベル
        for k in ['3連休抑制レベル', '3連休抑制', '3連休ペナルティ']:
            if k in raw_settings:
                val = str(raw_settings[k]).strip()
                if val in ['強', '中', '弱', 'なし']:
                    config.consecutive_off_level = val

        # 7. セクション人数ブレ幅
        for k in ['セクション人数ブレ幅', '人数ブレ幅', '許容ブレ幅']:
            if k in raw_settings:
                try: config.sec_capacity_tolerance = int(float(raw_settings[k]))
                except: pass

        # 8. 週休配分目標
        for k in ['週休配分目標', '週休目標', '週休配分']:
            if k in raw_settings:
                try: config.weekly_off_target = int(float(raw_settings[k]))
                except: pass

        # 9. 既存シフト維持レベル
        for k in ['既存シフト維持レベル', 'シフト維持レベル', '前回シフト維持']:
            if k in raw_settings:
                val = str(raw_settings[k]).strip()
                if val in ['強', '中', '弱']:
                    config.existing_shift_weight_level = val

        # 10. 2階裏 警告表示人数
        for k in ['2階裏_警告表示人数', '2階裏警告人数', '警告表示人数']:
            if k in raw_settings:
                try: config.warn_2f_back_count = int(float(raw_settings[k]))
                except: pass

        # 11. 祝日週の最低1日公休設定 (v1.5)
        for k in ['祝日週_公休設定', '祝日週公休設定', '祝日週_最低1日公休', '祝日週_固定公休振替', '祝日週固定公休振替', '固定公休振替', '祝日週振替公休']:
            if k in raw_settings:
                val = str(raw_settings[k]).strip()
                if val in ['無効', 'オフ', 'False', 'false', '0', 'しない', 'なし']:
                    config.holiday_week_off_rule = False
                elif val in ['有効', 'オン', 'True', 'true', '1', 'する', 'あり']:
                    config.holiday_week_off_rule = True

        # 12. 週単位の公休平準化 (v1.6)
        for k in ['週単位_公休平準化', '週単位公休平準化', '週休平準化', '週単位_週休平準化']:
            if k in raw_settings:
                val = str(raw_settings[k]).strip()
                if val in ['無効', 'オフ', 'False', 'false', '0', 'しない', 'なし']:
                    config.weekly_off_balance_rule = False
                elif val in ['有効', 'オン', 'True', 'true', '1', 'する', 'あり']:
                    config.weekly_off_balance_rule = True

        # 13. 2階裏 通常週の4名最低日数 (v1.6)
        for k in ['2階裏_週4名最低日数', '2階裏週4名最低日数', '2階裏_通常週4名最低日数', '2階裏_週4名日数']:
            if k in raw_settings:
                try: config.sec_2f_back_weekly_4_min = int(float(raw_settings[k]))
                except: pass

        # 14. 空欄スタッフの公休固定＆ローテーション (v1.6改善 / v1.7)
        for k in ['空欄スタッフ_公休固定ローテーション', '空欄スタッフ_公休ローテーション', '空欄公休ローテーション', '公休曜日固定ローテーション', '公休ローテーション', '空欄公休固定']:
            if k in raw_settings:
                val = str(raw_settings[k]).strip()
                if val in ['無効', 'オフ', 'False', 'false', '0', 'しない', 'なし']:
                    config.auto_rotate_off_rule = False
                elif val in ['有効', 'オン', 'True', 'true', '1', 'する', 'あり']:
                    config.auto_rotate_off_rule = True

        # 15. 指定公休日のHARD適用 (v1.8)
        for k in ['指定公休日_有効', '指定公休日有効', '指定公休_有効', '指定公休有効', '指定公休日_HARD適用', '指定公休_HARD適用', '指定公休日']:
            if k in raw_settings:
                val = str(raw_settings[k]).strip()
                if val in ['無効', 'オフ', 'False', 'false', '0', 'しない', 'なし']:
                    config.specified_off_rule = False
                elif val in ['有効', 'オン', 'True', 'true', '1', 'する', 'あり']:
                    config.specified_off_rule = True

        print(f"★ 『{target_sheet_name}』シートから動的パラメータを正常に読み込みました。")


    except Exception as e:
        print(f"警告: 設定シートのパース中にエラーが発生しました（初期値を使用します）: {e}")

    return config

def get_days_in_month(year: int, month: int) -> int:
    """指定年月の総日数を取得"""
    if month == 12:
        next_month = datetime.date(year + 1, 1, 1)
    else:
        next_month = datetime.date(year, month + 1, 1)
    return (next_month - datetime.timedelta(days=1)).day

def parse_date_list(val_str: Any, default_year: int = 2026, holidays_jp_dict: dict = None) -> Set[datetime.date]:
    """日付文字列、Excelシリアル値(int/float/np.number)、datetime型等を堅牢にパースし、祝日・日曜の有休を除外"""
    if pd.isna(val_str) or val_str is None:
        return set()
    
    # 1. 既に datetime.date や datetime.datetime の場合
    if isinstance(val_str, (datetime.date, datetime.datetime, pd.Timestamp)):
        d_obj = val_str.date() if isinstance(val_str, (datetime.datetime, pd.Timestamp)) else val_str
        if holidays_jp_dict and (d_obj.weekday() == 6 or d_obj in holidays_jp_dict):
            return set()
        return {d_obj}

    # 2. Excelシリアル値 (int / float / numpy数値型) の場合 (例: 46342 -> 2026-11-16)
    if isinstance(val_str, numbers.Number) and not isinstance(val_str, bool):
        try:
            f_val = float(val_str)
            if 30000 <= f_val <= 60000:
                d_obj = datetime.date(1899, 12, 30) + datetime.timedelta(days=int(f_val))
                if holidays_jp_dict and (d_obj.weekday() == 6 or d_obj in holidays_jp_dict):
                    return set()
                return {d_obj}
        except Exception:
            pass

    val_str = unicodedata.normalize('NFKC', str(val_str)).strip()
    if not val_str:
        return set()

    # 3. 文字列全体が5桁前後の数字（シリアル値文字列）の場合
    if re.match(r'^\d{5}(\.0+)?$', val_str):
        try:
            num = int(float(val_str))
            if 30000 <= num <= 60000:
                d_obj = datetime.date(1899, 12, 30) + datetime.timedelta(days=num)
                if holidays_jp_dict and (d_obj.weekday() == 6 or d_obj in holidays_jp_dict):
                    return set()
                return {d_obj}
        except Exception:
            pass

    res = set()
    # 時刻部分 (00:00:00 等) をトリム
    val_str = re.sub(r'\b\d{1,2}:\d{2}(:\d{2})?\b', '', val_str).strip()
    raw_parts = re.split(r'[,、;・\n\r]+', val_str)
    parts = []
    for rp in raw_parts:
        rp = rp.strip()
        if not rp:
            continue
        # 範囲指定を含まない空白区切りの複数日付に対応
        if ' ' in rp and not re.search(r'[-〜~]', rp):
            parts.extend([p for p in rp.split() if p.strip()])
        else:
            parts.append(rp)

    def _parse_single_date(s: str) -> datetime.date:
        s = s.strip()
        if not s: return None
        s = s.replace('年', '/').replace('月', '/').replace('日', '').replace('.', '/')

        # 5桁シリアル値文字列の考慮
        if re.match(r'^\d{5}$', s):
            try:
                num = int(s)
                if 30000 <= num <= 60000:
                    return datetime.date(1899, 12, 30) + datetime.timedelta(days=num)
            except Exception: pass

        m_short = re.match(r'^(\d{1,2})[-/](\d{1,2})$', s)
        if m_short:
            m, d = int(m_short.group(1)), int(m_short.group(2))
            try: return datetime.date(default_year, m, d)
            except Exception: return None
            
        m_full = re.match(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$', s)
        if m_full:
            y, m, d = int(m_full.group(1)), int(m_full.group(2)), int(m_full.group(3))
            try: return datetime.date(y, m, d)
            except Exception: return None
            
        try:
            dt = pd.to_datetime(s, errors='coerce')
            if pd.notna(dt):
                return dt.date()
        except Exception:
            pass
        return None

    for part in parts:
        part = part.strip()
        if not part: continue
        
        range_match = None
        m_full_range = re.match(r'^(\d{4}[-/]\d{1,2}[-/]\d{1,2})\s*[-〜~]\s*(\d{4}[-/]\d{1,2}[-/]\d{1,2})$', part)
        m_short_range = re.match(r'^(\d{1,2}[-/]\d{1,2})\s*[-〜~]\s*(\d{1,2}[-/]\d{1,2})$', part)
        
        if m_full_range:
            range_match = (m_full_range.group(1), m_full_range.group(2))
        elif m_short_range:
            range_match = (m_short_range.group(1), m_short_range.group(2))
        elif '〜' in part:
            sp = part.split('〜')
            if len(sp) == 2: range_match = (sp[0], sp[1])
        elif '~' in part:
            sp = part.split('~')
            if len(sp) == 2: range_match = (sp[0], sp[1])
        elif ' - ' in part:
            sp = part.split(' - ')
            if len(sp) == 2: range_match = (sp[0], sp[1])

        if range_match:
            d_start = _parse_single_date(range_match[0])
            d_end = _parse_single_date(range_match[1])
            if d_start and d_end:
                if d_start > d_end:
                    d_start, d_end = d_end, d_start
                cur = d_start
                while cur <= d_end:
                    if not (holidays_jp_dict and (cur.weekday() == 6 or cur in holidays_jp_dict)):
                        res.add(cur)
                    cur += datetime.timedelta(days=1)
                continue
                
        d_single = _parse_single_date(part)
        if d_single:
            if not (holidays_jp_dict and (d_single.weekday() == 6 or d_single in holidays_jp_dict)):
                res.add(d_single)
                
    return res

def parse_rule_target_dates(r_cond: str, all_dates: list, holidays_jp: set, ope_dict: dict) -> list:
    """制約条件の適用日をパース"""
    r_cond = str(r_cond).strip()
    wd_map = {"月": 0, "火": 1, "水": 2, "木": 3, "金": 4, "土": 5, "日": 6}
    t_dates = []
    
    quote_kw = None
    if '「' in r_cond and '」' in r_cond:
        quote_kw = r_cond.split('「')[1].split('」')[0].strip()
    
    for d in all_dates:
        if d.weekday() == 6 or d in holidays_jp: continue
        op_info = ope_dict.get(d, "")
        is_ope = bool(op_info)
        
        if r_cond == '全日':
            t_dates.append(d)
        elif r_cond in ['月曜', '火曜', '水曜', '木曜', '金曜', '土曜']:
            if d.weekday() == wd_map[r_cond[0]]: t_dates.append(d)
        elif '日曜日以外の手術のある日' in r_cond or r_cond == '手術のある日':
            if is_ope: t_dates.append(d)
        elif 'OPE' in r_cond or '手術' in r_cond:
            req_wd = None
            if r_cond[0] in wd_map and '曜' in r_cond:
                req_wd = wd_map[r_cond[0]]
            if req_wd is not None and d.weekday() != req_wd: continue
            
            if quote_kw:
                if is_ope and (quote_kw.lower() in op_info.lower()): t_dates.append(d)
            elif '硝子体' in r_cond or 'Vit' in r_cond:
                if is_ope and ('Vit' in op_info or '硝子体' in op_info): t_dates.append(d)
            elif '下垂' in r_cond or '眼瞼' in r_cond:
                if is_ope and ('下垂' in op_info or '眼瞼' in op_info): t_dates.append(d)
            elif '白内障' in r_cond or 'cat' in r_cond.lower():
                if is_ope and ('cat' in op_info.lower() or '白内障' in op_info): t_dates.append(d)
            elif '緑内障' in r_cond or 'gla' in r_cond.lower():
                if is_ope and ('gla' in op_info.lower() or '緑内障' in op_info): t_dates.append(d)
            else:
                if is_ope: t_dates.append(d)
        elif d.weekday() == wd_map.get(r_cond[0], -1):
            t_dates.append(d)
    return t_dates

def parse_previous_month_sheet(ws) -> Tuple[Dict[str, List[int]], Dict[str, int]]:
    """
    前月確定シフトシートの読み込み（縦型・横型両対応）。
    返り値: (tail_shifts_dict, main_offs_dict)
    - tail_shifts_dict: {スタッフ名: [直前7日間の出勤フラグ(1/0)]}
    - main_offs_dict: {スタッフ名: 前月の主公休曜日(0-5)}
    """
    tail_shifts: Dict[str, List[int]] = {}
    main_offs: Dict[str, int] = {}
    
    data = [[cell.value for cell in row] for row in ws.iter_rows()]
    if len(data) < 2:
        return tail_shifts, main_offs

    # --- 形式判定 ---
    # 縦型（行が日付、列がスタッフ名）: 行2〜3にスタッフ名があり、行3以降の1列目が日付
    # 横型（列が日付、行がスタッフ名）: 1列目にスタッフ名があり、ヘッダーに日付が並ぶ
    
    # パターンA: 縦型判定（現場出力フォーマット）
    # 行1: タイトル, 行2: ヘッダー [日付, 曜日, 備考, 警告, スタッフ名...]
    header_row_idx = None
    for r_idx in range(min(5, len(data))):
        row_vals = [str(v).strip() for v in data[r_idx] if v is not None]
        if any('日付' in v for v in row_vals) and any('曜日' in v for v in row_vals):
            header_row_idx = r_idx
            break

    if header_row_idx is not None:
        # 縦型フォーマット
        headers = [str(c).strip() if c is not None else "" for c in data[header_row_idx]]
        # スタッフ列を特定 (列5以降などで「人数」「総出勤」を含まない列)
        staff_cols = {}
        for c_idx, h_name in enumerate(headers):
            if c_idx >= 4 and h_name and not any(k in h_name for k in ['人数', '総出勤', '合計', '所定', '公休', '曜日']):
                staff_cols[h_name] = c_idx

        # 日付行を収集
        date_rows = []
        summary_rows = {}
        wd_map = {"月": 0, "火": 1, "水": 2, "木": 3, "金": 4, "土": 5, "日": 6}

        for r_idx in range(header_row_idx + 1, len(data)):
            row = data[r_idx]
            d_val = row[0] if len(row) > 0 else None
            lbl_val = row[2] if len(row) > 2 else None  # 備考列に集計行ラベルが入る場合あり

            if str(lbl_val).strip() == "主公休曜日":
                for s_name, c_idx in staff_cols.items():
                    if c_idx < len(row):
                        v = str(row[c_idx]).strip()
                        if v in wd_map:
                            main_offs[s_name] = wd_map[v]
                continue

            if d_val is not None:
                d_str = str(d_val).strip()
                # 日付判定 (例: '10/01', '2026-10-01', datetime.date)
                if isinstance(d_val, (datetime.date, datetime.datetime)) or ('/' in d_str and len(d_str.split('/')) == 2):
                    date_rows.append(row)

        # 末尾7日分を取得
        tail_7_rows = date_rows[-7:] if len(date_rows) >= 7 else date_rows
        for s_name, c_idx in staff_cols.items():
            hist = []
            for row in tail_7_rows:
                v = str(row[c_idx]).strip() if c_idx < len(row) and row[c_idx] is not None else ""
                hist.append(1 if v == '○' else 0)
            tail_shifts[s_name] = hist

        if tail_shifts:
            print(f"前月実績シート (縦型フォーマット): {len(tail_shifts)}名の末尾{len(tail_7_rows)}日分実績を取得しました。")
            return tail_shifts, main_offs

    # パターンB: 横型フォーマット（1列目が氏名、列ヘッダーが日付）
    df_prev = pd.DataFrame(data[1:], columns=data[0])
    name_col = next((c for c in df_prev.columns if c and any(k in str(c) for k in ['氏名', '名前', 'スタッフ'])), df_prev.columns[0])
    date_cols = [c for c in df_prev.columns if c and (isinstance(c, (datetime.datetime, datetime.date, int)) or (isinstance(c, str) and '/' in c))]
    
    if date_cols:
        tail_cols = date_cols[-7:] if len(date_cols) >= 7 else date_cols
        for _, row in df_prev.iterrows():
            p_name = str(row.get(name_col, '')).strip()
            if not p_name or p_name == 'None': continue
            hist = []
            for c in tail_cols:
                val = str(row.get(c, ''))
                hist.append(1 if val == '○' else 0)
            tail_shifts[p_name] = hist
        print(f"前月実績シート (横型フォーマット): {len(tail_shifts)}名の末尾{len(tail_cols)}日分実績を取得しました。")

    return tail_shifts, main_offs

def load_and_parse_data(filepath: str) -> tuple:
    """Excelファイルから全シートを読み込み、最適化用データ構造を生成"""
    wb_in = openpyxl.load_workbook(filepath, data_only=True)
    
    # シート1: スタッフマスタ
    ws_staff = wb_in[wb_in.sheetnames[0]]
    data_staff = [[cell.value for cell in row] for row in ws_staff.iter_rows()]
    df_staff = pd.DataFrame(data_staff[1:], columns=data_staff[0])
    
    excl_col = next((c for c in df_staff.columns if c and ('除外' in str(c))), None)
    if excl_col:
        df_staff = df_staff[~df_staff[excl_col].isin([True, 1, "1", "○"])]
        
    vac_col = next((c for c in df_staff.columns if c and ('有休' in str(c) and '特別' not in str(c))), None)
    sp_vac_col = next((c for c in df_staff.columns if c and ('特別' in str(c))), None)
    pref_off_col = next((c for c in df_staff.columns if c and any(k in str(c) for k in ['希望休'])), None)
    specified_off_col = next((c for c in df_staff.columns if c and any(k in str(c) for k in ['指定公休', '公休指定'])), None)
    if not specified_off_col:
        specified_off_col = next((c for c in df_staff.columns if c and ('公休' in str(c) and '曜日' not in str(c) and any(k in str(c) for k in ['指定', '日']))), None)
    
    holidays_set = fetch_holidays()
    holidays_jp_dict = {d: name for d, name in holidays_set}
    
    staff_list = []
    for _, row in df_staff.iterrows():
        name_col = next((c for c in df_staff.columns if c and any(k in str(c) for k in ['氏名', '名前', 'スタッフ'])), df_staff.columns[0])
        name = str(row.get(name_col, '')).strip()
        if not name or name == 'None': continue
        
        sec = str(row.get('所属', '')).strip()
        job = str(row.get('職種', '')).strip()
        emp = str(row.get('雇用形態', '')).strip()
        pat = str(row.get('基本勤務パターン', '')).strip()
        pref = str(row.get('希望出勤曜日', '')).strip()
        
        is_core = True if emp in ['正社員', 'パート社員'] else False
        
        s_info = {
            'name': name,
            'section': sec,
            'job_role': job,
            'emp_type': emp,
            'pattern': pat,
            'is_core': is_core,
            'is_part_1f': (sec == '1階' and emp == 'パート'), 
            'pref_weekdays_str': pref,
            'raw_vacation': row.get(vac_col),
            'raw_special_vacation': row.get(sp_vac_col),
            'raw_preferred_off': row.get(pref_off_col),
            'raw_specified_off': row.get(specified_off_col)
        }
        
        wd_map = {"月": 0, "火": 1, "水": 2, "木": 3, "金": 4, "土": 5, "日": 6}
        boff_list = []
        boff_str = str(row.get('希望公休曜日', row.get('基本公休', ''))).strip()
        for char in boff_str:
            if char in wd_map and wd_map[char] not in boff_list:
                boff_list.append(wd_map[char])

        s_info['basic_off_days'] = boff_list

        if 'もしくは' in boff_str or 'あるいは' in boff_str:
            s_info['basic_off_type'] = 'OR'
        elif boff_list:
            s_info['basic_off_type'] = 'AND'
        else:
            s_info['basic_off_type'] = 'OR'
            
        staff_list.append(s_info)

    # シート2: シフト制約条件
    ws_rules = wb_in[wb_in.sheetnames[1]]
    data_rules = [[cell.value for cell in row] for row in ws_rules.iter_rows()]
    df_rules = pd.DataFrame(data_rules[1:], columns=data_rules[0])
    
    # 【小西眼科など外部発生イベントの安全な除外】
    cond_col = next((c for c in df_rules.columns if c and any(k in str(c) for k in ['条件', '適用'])), df_rules.columns[1])
    if cond_col:
        df_rules = df_rules[~df_rules[cond_col].astype(str).str.contains('小西', na=False)].reset_index(drop=True)

    # シート3: 手術スケジュール
    ws_ope = wb_in[wb_in.sheetnames[2]]
    data_ope = [[cell.value for cell in row] for row in ws_ope.iter_rows()]
    df_ope = pd.DataFrame(data_ope[1:], columns=data_ope[0])
    
    df_ope['日付_dt'] = pd.to_datetime(df_ope['日付']).dt.date
    col_type = next((c for c in df_ope.columns if c and any(k in str(c) for k in ['種別', '内容', '名称', '手術', 'OPE'])), df_ope.columns[1])
    
    ope_dict_all = {}
    for _, row in df_ope.iterrows():
        if pd.notna(row['日付_dt']):
            ope_dict_all[row['日付_dt']] = str(row.get(col_type, '')).strip()

    all_months = set()
    tgt_col = next((c for c in df_ope.columns if c and '作成対象' in str(c)), None)
    for _, row in df_ope.iterrows():
        if pd.notna(row['日付_dt']):
            all_months.add((row['日付_dt'].year, row['日付_dt'].month))
                
    target_months = sorted(list(all_months))

    # シート4: 前月シフト実績（存在する場合）
    prev_month_tail_shifts = {}
    prev_month_main_offs = {}
    if len(wb_in.sheetnames) > 3:
        try:
            ws_prev = wb_in[wb_in.sheetnames[3]]
            prev_month_tail_shifts, prev_month_main_offs = parse_previous_month_sheet(ws_prev)
        except Exception as e:
            print(f"前月実績シートの解析警告: {e}")

    # 設定シートの読み込み
    config = parse_settings_sheet(wb_in)

    # 空欄スタッフへの公休曜日月内固定＆ローテーションルールの適用 (v1.6改善 / v1.7)
    if config.auto_rotate_off_rule:
        for s in staff_list:
            if s.get('is_core') and s.get('pattern') != '除外' and not s.get('is_part_1f'):
                if not s.get('basic_off_days'):
                    # 希望公休曜日が空欄のスタッフに平日(月〜金)の主公休ローテーションルールを適用
                    s['basic_off_days'] = [0, 1, 2, 3, 4]
                    s['basic_off_type'] = 'OR'
                    s['is_auto_rotated_off'] = True

    wb_in.close()
    return staff_list, df_rules, ope_dict_all, target_months, holidays_jp_dict, prev_month_tail_shifts, prev_month_main_offs, config


def inspect_excel_file(filepath: str) -> dict:
    """
    GUIプレビュー用にExcelファイルのメタ情報を軽量解析
    """
    info = {
        'filepath': filepath,
        'sheet_names': [],
        'target_months': [],
        'staff_count': 0,
        'has_prev_month': False,
        'prev_month_sheet_name': None,
        'has_settings': False,
        'settings_sheet_name': None,
        'rules_count': 0
    }
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        info['sheet_names'] = wb.sheetnames
        
        # スタッフマスタ
        if len(wb.sheetnames) >= 1:
            ws_staff = wb[wb.sheetnames[0]]
            info['staff_count'] = max(0, ws_staff.max_row - 1)
            
        # シフト制約条件
        if len(wb.sheetnames) >= 2:
            ws_rules = wb[wb.sheetnames[1]]
            info['rules_count'] = max(0, ws_rules.max_row - 1)
            
        # 手術スケジュール
        if len(wb.sheetnames) >= 3:
            ws_ope = wb[wb.sheetnames[2]]
            data_ope = [[cell.value for cell in row] for row in ws_ope.iter_rows()]
            if len(data_ope) > 1:
                df_ope = pd.DataFrame(data_ope[1:], columns=data_ope[0])
                if '日付' in df_ope.columns:
                    df_ope['日付_dt'] = pd.to_datetime(df_ope['日付'], errors='coerce').dt.date
                    tgt_col = next((c for c in df_ope.columns if c and '作成対象' in str(c)), None)
                    all_months = set()
                    marked_months = set()
                    for _, row in df_ope.iterrows():
                        if pd.notna(row['日付_dt']):
                            y = row['日付_dt'].year
                            m = row['日付_dt'].month
                            all_months.add((y, m))
                            if tgt_col and row.get(tgt_col) in ['○', 1, True, "1"]:
                                marked_months.add((y, m))
                    info['target_months'] = sorted(list(all_months))
                    info['marked_months'] = sorted(list(marked_months))

        # 前月実績
        for s in wb.sheetnames:
            if any(k in s for k in ['実績', '前月']):
                info['has_prev_month'] = True
                info['prev_month_sheet_name'] = s
                break

        # 設定シート
        for s in wb.sheetnames:
            if any(k in s for k in ['設定', 'システム設定', 'パラメータ', 'config']):
                info['has_settings'] = True
                info['settings_sheet_name'] = s
                break

        wb.close()
    except Exception as e:
        info['error'] = str(e)

    return info
