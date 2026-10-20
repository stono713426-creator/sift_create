"""
Excel描画モジュール (openpyxl)
現場フォーマットの完全再現（印刷設定、列幅、行高さ、フォント、罫線、数式埋め込み）
"""
import datetime
from typing import Dict, Tuple, Optional, Any
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .config import SECTION_DEFINITIONS, ShiftConfig
from .parser import get_days_in_month

def apply_cell_style(
    cell: openpyxl.cell.cell.Cell,
    value: Any = None,
    font: Optional[Font] = None,
    fill: Optional[PatternFill] = None,
    border: Optional[Border] = None,
    alignment: Optional[Alignment] = None
) -> None:
    """セルへのスタイル一括適用ヘルパー"""
    if value is not None: cell.value = value
    if font: cell.font = font
    if fill: cell.fill = fill
    if border: cell.border = border
    if alignment: cell.alignment = alignment

def render_monthly_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    best_x: Optional[Dict[Tuple[str, datetime.date], int]],
    best_y: Optional[Dict[str, int]],
    monthly_data: dict,
    target_days_dict: Optional[Dict[str, int]]
) -> None:
    """単月のシフト表を現場フォーマットで描画"""
    year = monthly_data['year']
    month = monthly_data['month']
    
    if best_x is None:
        ws.title = f"{year}年{month:02d}月_未生成"
        f_err = Font(name="メイリオ", size=12, bold=True, color="C00000")
        apply_cell_style(ws['A1'], f"【制約条件の競合により生成不可】 {year}年{month}月のシフト表は条件が厳しすぎるため生成できませんでした。", font=f_err)
        return

    staff_list = monthly_data['staff_list']
    ope_dict = monthly_data['ope_dict']
    holidays_jp_dict = monthly_data['holidays_jp_dict']
    cfg: ShiftConfig = monthly_data.get('config') or ShiftConfig()
    
    ws.title = f"{year}年{month:02d}月シフト"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "E3"
    
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    
    days_in_month = get_days_in_month(year, month)
    all_dates = [datetime.date(year, month, d) for d in range(1, days_in_month + 1)]
    
    sections = SECTION_DEFINITIONS
    
    f_title = Font(name="メイリオ", size=14, bold=True, color="000000")
    f_white = Font(name="メイリオ", size=10, bold=True, color="FFFFFF")
    
    fill_holiday = PatternFill("solid", fgColor="FCE4D6")
    fill_warn = PatternFill("solid", fgColor="FCE4D6")
    fill_none = PatternFill("none")
    
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    border_double_top = Border(
        top=Side(style='double', color='000000'), left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    
    align_c = Alignment(horizontal="center", vertical="center")
    align_l = Alignment(horizontal="left", vertical="center")

    ws.merge_cells('A1:K1')
    apply_cell_style(ws['A1'], f"{year}年{month}月 勤務シフト表", font=f_title, alignment=align_l)
    ws.row_dimensions[1].height = 30.0
    
    fixed_headers = [("日付", 6.25), ("曜日", 5.0), ("備考", 11.625), ("警告", 5.0)]
    for c, (txt, w) in enumerate(fixed_headers, 1):
        col_let = get_column_letter(c)
        ws.column_dimensions[col_let].width = w
        apply_cell_style(ws.cell(row=2, column=c), txt, font=f_white, fill=PatternFill("solid", fgColor="305496" if c < 4 else "C00000"), alignment=align_c, border=border_thin)

    col_idx = 5
    sec_col_ranges = {}
    staff_col_map = {}
    
    for sec_info in sections:
        sec = sec_info["id"]
        s_fill = PatternFill("solid", fgColor=sec_info["staff_fill"])
        c_fill = PatternFill("solid", fgColor=sec_info["cnt_fill"])
        
        sec_staffs = [s for s in staff_list if s['section'] == sec]
        start_col = col_idx
        
        for s in sec_staffs:
            col_let = get_column_letter(col_idx)
            ws.column_dimensions[col_let].width = 13.0
            apply_cell_style(ws.cell(row=2, column=col_idx), s['name'], font=f_white, fill=s_fill, alignment=align_c, border=border_thin)
            staff_col_map[s['name']] = col_idx
            col_idx += 1
            
        end_col = col_idx - 1
        
        cnt_col = col_idx
        col_let = get_column_letter(cnt_col)
        ws.column_dimensions[col_let].width = sec_info["cnt_width"]
        apply_cell_style(ws.cell(row=2, column=cnt_col), f"{sec} 人数", font=f_white, fill=c_fill, alignment=align_c, border=border_thin)
        
        sec_col_ranges[sec] = (start_col, end_col, cnt_col)
        col_idx += 1
        
        if sec_info["blank"]:
            col_let = get_column_letter(col_idx)
            ws.column_dimensions[col_let].width = 4.0
            col_idx += 1
            
    total_cnt_col = col_idx
    col_let = get_column_letter(total_cnt_col)
    ws.column_dimensions[col_let].width = 9.875
    apply_cell_style(ws.cell(row=2, column=total_cnt_col), "総出勤人数", font=f_white, fill=PatternFill("solid", fgColor="1B365D"), alignment=align_c, border=border_thin)
    
    ws.row_dimensions[2].height = 28.0
    
    wd_str = ["月", "火", "水", "木", "金", "土", "日"]
    for r, d in enumerate(all_dates, 3):
        ws.row_dimensions[r].height = 22.0
        
        is_holiday = (d.weekday() == 6) or (d in holidays_jp_dict)
        has_ope = bool(ope_dict.get(d, ""))
        
        remark_texts = []
        if d in holidays_jp_dict: remark_texts.append(holidays_jp_dict[d])
        if has_ope: remark_texts.append(ope_dict[d])
        remark_str = " / ".join(remark_texts) if remark_texts else ""
        
        row_fill = fill_holiday if is_holiday else fill_none
        f_row_data = Font(name="メイリオ", size=10, color="C00000" if is_holiday else "000000")
        f_gray = Font(name="メイリオ", size=10, color="595959")
        f_green = Font(name="メイリオ", size=10, color="006100")
        f_red = Font(name="メイリオ", size=10, color="C00000")
        
        apply_cell_style(ws.cell(row=r, column=1), d.strftime("%m/%d"), font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
        apply_cell_style(ws.cell(row=r, column=2), wd_str[d.weekday()], font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
        apply_cell_style(ws.cell(row=r, column=3), remark_str, font=f_row_data, fill=row_fill, alignment=align_l, border=border_thin)
        
        c_d = ws.cell(row=r, column=4)
        apply_cell_style(c_d, "", fill=row_fill, alignment=align_c, border=border_thin)
        
        cnt_2f_b = 0
        for s in staff_list:
            sname = s['name']
            c_idx = staff_col_map[sname]
            c_cell = ws.cell(row=r, column=c_idx)
            
            if is_holiday:
                apply_cell_style(c_cell, "休業", font=f_red, fill=row_fill, alignment=align_c, border=border_thin)
                continue
                
            if d in s.get('vacation_dates', set()):
                apply_cell_style(c_cell, "有休", font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
            elif d in s.get('special_vacation_dates', set()):
                apply_cell_style(c_cell, "特有", font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
            else:
                is_work = best_x[(sname, d)] == 1
                if is_work:
                    apply_cell_style(c_cell, "○", font=f_green if has_ope else f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
                    if s['section'] == '2階裏': cnt_2f_b += 1
                else:
                    apply_cell_style(c_cell, "公", font=f_gray, fill=row_fill, alignment=align_c, border=border_thin)
                    
        if not is_holiday and cnt_2f_b <= cfg.warn_2f_back_count:
            apply_cell_style(c_d, "⚠️", font=f_red, fill=fill_warn, alignment=align_c, border=border_thin)
        
        for sec, (sc, ec, cc) in sec_col_ranges.items():
            c_cell = ws.cell(row=r, column=cc)
            val = ""
            if not is_holiday and sc <= ec:
                l_sc = get_column_letter(sc)
                l_ec = get_column_letter(ec)
                val = f'=COUNTIF({l_sc}{r}:{l_ec}{r}, "○")'
            apply_cell_style(c_cell, val, font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
            
        c_cell = ws.cell(row=r, column=total_cnt_col)
        val = ""
        if not is_holiday:
            sum_parts = []
            for _, (_, _, cc) in sec_col_ranges.items():
                sum_parts.append(f"{get_column_letter(cc)}{r}")
            if sum_parts:
                val = f'=SUM({",".join(sum_parts)})'
        apply_cell_style(c_cell, val, font=f_row_data, fill=row_fill, alignment=align_c, border=border_thin)
        
    end_row = 2 + len(all_dates)
    sum_rows = [end_row + i for i in range(1, 5)]
    for sr in sum_rows: ws.row_dimensions[sr].height = 24.0
    
    labels = ["合計出勤日数", "所定勤務日数", "公休日数", "主公休曜日"]
    f_sum = Font(name="メイリオ", size=10, bold=True, color="000000")
    for i, lbl in enumerate(labels):
        apply_cell_style(ws.cell(row=sum_rows[i], column=3), lbl, font=f_sum, alignment=align_c, border=border_double_top if i == 0 else border_thin)
        
    for s in staff_list:
        c_idx = staff_col_map[s['name']]
        c_let = get_column_letter(c_idx)
        
        apply_cell_style(ws.cell(row=sum_rows[0], column=c_idx), f'=COUNTIF({c_let}3:{c_let}{end_row}, "○")', font=f_sum, alignment=align_c, border=border_double_top)
        apply_cell_style(ws.cell(row=sum_rows[1], column=c_idx), target_days_dict[s['name']], font=f_sum, alignment=align_c, border=border_thin)
        apply_cell_style(ws.cell(row=sum_rows[2], column=c_idx), f'=COUNTIF({c_let}3:{c_let}{end_row}, "公")', font=f_sum, alignment=align_c, border=border_thin)
        
        m_off = best_y.get(s['name'], -1)
        has_rot_2nd = (
            cfg.auto_rotate_off_rule
            and s.get('is_core', False)
            and ('週4' in s.get('pattern', ''))
            and s.get('basic_off_type') == 'AND'
            and len(s.get('basic_off_days', [])) == 1
        )
        if has_rot_2nd and m_off >= 0:
            fixed_w = s['basic_off_days'][0]
            val = f"{wd_str[fixed_w]}・{wd_str[m_off]}" if fixed_w != m_off else wd_str[fixed_w]
        else:
            val = wd_str[m_off] if m_off >= 0 else "-"
        apply_cell_style(ws.cell(row=sum_rows[3], column=c_idx), val, font=f_sum, alignment=align_c, border=border_thin)
