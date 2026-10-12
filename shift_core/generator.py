"""
シフト生成統合サービスモジュール (GUI / CLI 両対応アーキテクチャ)
コールバックによる進捗通知、月またぎ連動、結果構造体の返却をサポート。
"""
import os
import glob
from dataclasses import dataclass, field
from typing import List, Dict, Tuple, Optional, Callable, Any
import openpyxl

from .parser import load_and_parse_data
from .solver import solve_monthly_shift
from .renderer import render_monthly_sheet

@dataclass
class MonthResult:
    year: int
    month: int
    success: bool
    stage: int
    message: str

@dataclass
class GenerationResult:
    success: bool
    output_filepath: Optional[str] = None
    month_results: List[MonthResult] = field(default_factory=list)
    error_message: Optional[str] = None

class ShiftGeneratorService:
    """勤務シフト自動生成サービスクラス"""

    def __init__(self, logger: Optional[Callable[[str], None]] = None):
        self.logger = logger or print

    def log(self, message: str) -> None:
        self.logger(message)

    def generate(
        self,
        input_file: Optional[str] = None,
        target_months: Optional[List[Tuple[int, int]]] = None,
        output_file: Optional[str] = None,
        on_progress: Optional[Callable[[float, str], None]] = None
    ) -> GenerationResult:
        """
        シフト自動生成を実行
        :param input_file: 入力Excelファイルのパス (Noneの場合はカレントディレクトリから最新を検索)
        :param target_months: 生成対象年月リスト [(2026, 10), ...] (Noneの場合はExcelから自動判定)
        :param output_file: 出力Excelファイルのパス (Noneの場合は自動生成)
        :param on_progress: 進捗通知コールバック func(progress_ratio: float, status_text: str)
        """
        def report_progress(ratio: float, text: str):
            if on_progress:
                try:
                    on_progress(ratio, text)
                except Exception:
                    pass
            self.log(f"[{int(ratio * 100)}%] {text}")

        report_progress(0.05, "入力ファイルの検索・読み込み中...")

        # 入力ファイルの特定
        if not input_file:
            files = glob.glob('*勤務条件表*.xlsx') + glob.glob('ope勤務条件表*.xlsx')
            files = list(set(files))
            if not files:
                err = "カレントディレクトリに対象の勤務条件表Excelファイルが見つかりません。"
                self.log(f"エラー: {err}")
                return GenerationResult(success=False, error_message=err)
            input_file = max(files, key=os.path.getmtime)

        self.log(f"入力ファイル: {input_file}")

        try:
            # データパース
            (
                staff_list,
                df_rules,
                ope_dict_all,
                auto_target_months,
                holidays_jp_dict,
                init_tail_shifts,
                init_main_offs,
                config
            ) = load_and_parse_data(input_file)
        except Exception as e:
            err = f"データ読み込みエラー: {e}"
            self.log(err)
            return GenerationResult(success=False, error_message=err)

        # 対象年月の決定
        months_to_process = target_months if target_months else auto_target_months
        if not months_to_process:
            err = "作成対象の月が指定されていません（手術スケジュールの『作成対象』列を確認してください）。"
            self.log(f"エラー: {err}")
            return GenerationResult(success=False, error_message=err)

        report_progress(0.15, f"対象月 ({len(months_to_process)}ヶ月) のモデル構築を開始します...")

        wb_out = openpyxl.Workbook()
        if "Sheet" in wb_out.sheetnames:
            del wb_out["Sheet"]

        # 前月実績データの初期化（第4シートから読み込んだ実績、または空）
        prev_tail_shifts = init_tail_shifts
        prev_main_offs = init_main_offs

        if prev_tail_shifts:
            self.log(f"★ 前月実績シートから {len(prev_tail_shifts)} 名の実績データを最初の対象月に連動させます。")
        else:
            self.log("前月実績データなし（単月初期状態として計算を開始します）。")

        month_results: List[MonthResult] = []
        total_months = len(months_to_process)

        for idx, (t_year, t_month) in enumerate(months_to_process):
            base_ratio = 0.15 + (idx / total_months) * 0.75
            month_label = f"{t_year}年{t_month:02d}月"
            report_progress(base_ratio, f"【{month_label}】の最適化計算中...")
            
            self.log(f"\n==========================================")
            self.log(f"  対象月処理開始: {month_label}")
            if prev_tail_shifts:
                self.log(f"  -> 前月実績連動: {len(prev_tail_shifts)}名の直前勤務データを引き継ぎ中")
            self.log(f"==========================================")

            ope_dict_month = {d: v for d, v in ope_dict_all.items() if d.year == t_year and d.month == t_month}
            monthly_data = {
                'staff_list': staff_list,
                'df_rules': df_rules,
                'ope_dict': ope_dict_month,
                'holidays_jp_dict': holidays_jp_dict,
                'target_file': input_file,
                'year': t_year,
                'month': t_month,
                'config': config
            }
            prev_m_data = {
                'tail_shifts': prev_tail_shifts,
                'main_offs': prev_main_offs
            }

            best_x, best_y, t_days, new_tail = None, None, None, None
            solved_stage = 0

            for stage in [1, 2]:
                self.log(f"Stage {stage} 最適化開始...")
                b_x, b_y, t_d, n_t = solve_monthly_shift(t_year, t_month, monthly_data, prev_m_data, stage)
                if b_x is not None:
                    best_x, best_y, t_days, new_tail = b_x, b_y, t_d, n_t
                    solved_stage = stage
                    self.log(f"Stage {stage} 成功.")
                    break

            ws = wb_out.create_sheet()
            if best_x is not None:
                self.log(f"【成功】{month_label}のシフトが生成されました。")
                # 次の月へ末尾実績と主公休を引き継ぐ（月跨ぎ連動）
                prev_tail_shifts = new_tail
                prev_main_offs = best_y
                render_monthly_sheet(ws, best_x, best_y, monthly_data, t_days)
                month_results.append(MonthResult(
                    year=t_year,
                    month=t_month,
                    success=True,
                    stage=solved_stage,
                    message=f"Stage {solved_stage} にて解を正常生成"
                ))
            else:
                self.log(f"【注意】{month_label}は条件の競合によりシフトを作成できませんでした。(INFEASIBLE)")
                render_monthly_sheet(ws, None, None, monthly_data, None)
                month_results.append(MonthResult(
                    year=t_year,
                    month=t_month,
                    success=False,
                    stage=0,
                    message="Stage 1/2ともに解なし (INFEASIBLE)"
                ))

        report_progress(0.92, "Excelブックの保存処理中...")

        # 出力ファイル名
        if not output_file:
            start_y, start_m = months_to_process[0]
            end_y, end_m = months_to_process[-1]
            output_file = f"{str(start_y)[-2:]}年{start_m:02d}-{end_m:02d}月シフト.xlsx"

        try:
            wb_out.save(output_file)
            wb_out.close()
            self.log(f"\nすべての対象月の出力処理が正常終了しました: {output_file}")
            report_progress(1.0, f"完了: {output_file} を出力しました。")
            return GenerationResult(
                success=True,
                output_filepath=output_file,
                month_results=month_results
            )
        except Exception as e:
            err = f"Excelファイル保存エラー: {e}"
            self.log(err)
            return GenerationResult(success=False, error_message=err, month_results=month_results)
