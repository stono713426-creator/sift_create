"""
数理最適化ソルバーモジュール (Google OR-Tools CP-SAT)
二段階最適化 (Stage 1: 厳格探索, Stage 2: スラック緩和探索)
"""
import re
import math
import datetime
from typing import Dict, List, Tuple, Optional, Any
import pandas as pd
from ortools.sat.python import cp_model
import openpyxl

from .config import (
    LEVEL_WEIGHT_MAP,
    REC_WEIGHT_MAP,
    WEIGHT_FIXED_OFF_SHORTAGE,
    WEIGHT_2F_BACK_MIN_SLACK,
    WEIGHT_GENERAL_SEC_SLACK,
    WEIGHT_NURSE_DEFICIT_SLACK,
    WEIGHT_SUBSTITUTE_OFF_SLACK,
    WEIGHT_2F_BACK_5_PENALTY,
    WEIGHT_SEC_BALANCE_SQ,
    WEIGHT_SEC_BALANCE_ABS,
    WEIGHT_MAIN_OFF_MISMATCH,
    WEIGHT_2F_BACK_3_PENALTY,
    WEIGHT_NURSE_LT2_PENALTY,
    WEIGHT_SPECIAL_EXTRA_WORK,
    WEIGHT_ROLLING_WINDOW_DIFF,
    WEIGHT_WEEKLY_OFF_DIFF,
    WEIGHT_2F_BACK_WEEKLY_4_SLACK,
    WEIGHT_PREV_OFF_REPEAT,
    WEIGHT_TAIL_OFF_SUM,
    SOLVER_MAX_TIME_SECONDS,
    SOLVER_NUM_SEARCH_WORKERS,
    ShiftConfig,
)
from .parser import get_days_in_month, parse_date_list, parse_rule_target_dates

def solve_monthly_shift(
    target_year: int,
    target_month: int,
    monthly_data: dict,
    prev_month_data: dict,
    stage: int
) -> tuple:
    """単月のシフト最適化モデルを構築・求解"""
    cfg: ShiftConfig = monthly_data.get('config') or ShiftConfig()
    staff_list = monthly_data['staff_list']
    df_rules = monthly_data['df_rules']
    ope_dict = monthly_data['ope_dict']
    holidays_jp_dict = monthly_data['holidays_jp_dict']
    target_file = monthly_data.get('target_file')
    
    prev_month_tail_shifts = prev_month_data.get('tail_shifts', {})
    prev_month_main_offs = prev_month_data.get('main_offs', {})

    days_in_month = get_days_in_month(target_year, target_month)
    all_dates = [datetime.date(target_year, target_month, d) for d in range(1, days_in_month + 1)]
    holidays_jp = set(holidays_jp_dict.keys())
    
    cal_days = len(all_dates)
    sundays = sum(1 for d in all_dates if d.weekday() == 6)
    holidays_in_month = sum(1 for d in all_dates if d in holidays_jp and d.weekday() != 6)
    
    for s in staff_list:
        s['vacation_dates'] = parse_date_list(s['raw_vacation'], target_year, holidays_jp_dict)
        s['special_vacation_dates'] = parse_date_list(s['raw_special_vacation'], target_year, holidays_jp_dict)
        s['preferred_off_dates'] = parse_date_list(s['raw_preferred_off'], target_year)
        s['specified_off_dates'] = parse_date_list(s.get('raw_specified_off'), target_year, holidays_jp_dict) if cfg.specified_off_rule else set()

    part_schedule_dict = {}
    wd_map = {"月": 0, "火": 1, "水": 2, "木": 3, "金": 4, "土": 5, "日": 6}
    for s in staff_list:
        if s.get('is_part_1f', False) or (not s.get('is_core')):
            pref_w_str = s.get('pref_weekdays_str', '')
            pref_w = [wd_map[w] for w in wd_map if w in pref_w_str]
            s['pref_weekdays'] = pref_w
            for d in all_dates:
                if d.weekday() in pref_w and d.weekday() != 6 and d not in holidays_jp:
                    vac_set = s.get('vacation_dates', set()) | s.get('special_vacation_dates', set()) | s.get('preferred_off_dates', set()) | s.get('specified_off_dates', set())
                    if d not in vac_set:
                        part_schedule_dict[(s['name'], d)] = 1
                    else:
                        part_schedule_dict[(s['name'], d)] = 0
                else:
                    part_schedule_dict[(s['name'], d)] = 0

    staff_col_idx = {}
    c_idx = 5
    for sec in ["2階表", "2階裏", "1階"]:
        sec_staffs = [s for s in staff_list if s['section'] == sec]
        for s in sec_staffs:
            staff_col_idx[s['name']] = c_idx
            c_idx += 1
        c_idx += 1 
        if sec != "1階":
            c_idx += 1

    target_days_dict = {}
    sec_target_sum = {"1階": 0, "2階表": 0, "2階裏": 0}
    
    for s in staff_list:
        name = s['name']
        pattern = s.get('pattern', '')
        vac_count = sum(1 for d in all_dates if d in s.get('vacation_dates', set()))
        sp_vac_count = sum(1 for d in all_dates if d in s.get('special_vacation_dates', set()))
        spec_off_count = sum(1 for d in all_dates if d in s.get('specified_off_dates', set()) and d.weekday() != 6 and d not in holidays_jp)
        
        if not s.get('is_core'):
            pref_w = s.get('pref_weekdays', [])
            vac_set = s.get('vacation_dates', set()) | s.get('special_vacation_dates', set()) | s.get('preferred_off_dates', set()) | s.get('specified_off_dates', set())
            target_days = sum(1 for d in all_dates if d.weekday() in pref_w and d.weekday() != 6 and d not in holidays_jp and d not in vac_set)
        else:
            basic_slots = cal_days - (sundays + holidays_in_month)
            wd_counts = {w: sum(1 for d in all_dates if d.weekday() == w) for w in range(6)}
            
            if '週4' in pattern:
                off_days = wd_counts[5] + wd_counts[4]
            else:
                off_days = wd_counts[5]
            
            effective_off_days = max(off_days, spec_off_count)
            target_days = max(0, basic_slots - effective_off_days - vac_count - sp_vac_count)

        target_days_dict[name] = target_days
        sec = s.get('section', '不明')
        # 平準化のための M_s: 1階はパートも含めた全体で合算、2階はコア要員のみ合算
        if sec == "1階":
            if sec in sec_target_sum:
                sec_target_sum[sec] += target_days
        elif sec in sec_target_sum and s.get('is_core'):
            sec_target_sum[sec] += target_days

    work_days = cal_days - sundays - holidays_in_month
    ideal_mean = {sec: (t_sum / work_days if work_days > 0 else 0) for sec, t_sum in sec_target_sum.items()}

    # カレンダー週（ISO週）のグループ化
    weeks = {}
    for d in all_dates:
        iso_year, iso_week, _ = d.isocalendar()
        w_key = (iso_year, iso_week)
        if w_key not in weeks:
            weeks[w_key] = []
        weeks[w_key].append(d)

    model = cp_model.CpModel()
    penalties = []
    x = {}
    y_mo = {}
    
    existing_shift_dict = {}
    if target_file:
        sheet_name = f"{target_year}年{target_month:02d}月シフト"
        try:
            wb_ref = openpyxl.load_workbook(target_file, data_only=True)
            if sheet_name in wb_ref.sheetnames:
                ws_exist = wb_ref[sheet_name]
                for r in range(3, ws_exist.max_row + 1):
                    d_val = ws_exist.cell(row=r, column=1).value
                    if not d_val: continue
                    try:
                        if isinstance(d_val, datetime.date): curr_d = d_val
                        else:
                            m_s, d_s = map(int, str(d_val).split('/'))
                            curr_d = datetime.date(target_year, m_s, d_s)
                    except: continue
                    
                    for s in staff_list:
                        sname = s['name']
                        if sname in staff_col_idx:
                            c_v = ws_exist.cell(row=r, column=staff_col_idx[sname]).value
                            existing_shift_dict[(sname, curr_d)] = 1 if c_v == '○' else 0
            wb_ref.close()
        except: pass

    for s in staff_list:
        name = s['name']
        for d in all_dates:
            x[(name, d)] = model.NewBoolVar(f"x_{name}_{d.strftime('%Y%m%d')}")
        for w in range(6):
            y_mo[(name, w)] = model.NewBoolVar(f"y_mo_{name}_{w}")

    if existing_shift_dict:
        for s in staff_list:
            sname = s['name']
            for d in all_dates:
                if (sname, d) in existing_shift_dict:
                    prev_val = existing_shift_dict[(sname, d)]
                    if d in s['vacation_dates'] or d in s['special_vacation_dates']: continue
                    diff_var = model.NewBoolVar(f"shift_change_{sname}_{d}")
                    if prev_val == 1:
                        model.Add(x[(sname, d)] == 0).OnlyEnforceIf(diff_var)
                        model.Add(x[(sname, d)] == 1).OnlyEnforceIf(diff_var.Not())
                    else:
                        model.Add(x[(sname, d)] == 1).OnlyEnforceIf(diff_var)
                        model.Add(x[(sname, d)] == 0).OnlyEnforceIf(diff_var.Not())
                    penalties.append(diff_var * cfg.weight_shift_change)

    for s in staff_list:
        name = s['name']
        pattern = s.get('pattern', '')
        is_special = '必要に応じて週5' in pattern
        target_days = target_days_dict[name]
        
        basic_slots = cal_days - (sundays + holidays_in_month)
        vac_count = sum(1 for d in all_dates if d in s.get('vacation_dates', set()))
        sp_vac_count = sum(1 for d in all_dates if d in s.get('special_vacation_dates', set()))
        
        if not s.get('is_core'):
            for d in all_dates:
                model.Add(x[(name, d)] == part_schedule_dict[(name, d)])
        else:
            for d in all_dates:
                if d.weekday() == 6 or d in holidays_jp or d in s.get('vacation_dates', set()) or d in s.get('special_vacation_dates', set()) or d in s.get('preferred_off_dates', set()) or d in s.get('specified_off_dates', set()):
                    model.Add(x[(name, d)] == 0)

            month_sum_var = sum(x[(name, d)] for d in all_dates)

            if s.get('basic_off_type') == 'AND':
                for d in all_dates:
                    if d.weekday() in s.get('basic_off_days', []) and d not in holidays_jp:
                        model.Add(x[(name, d)] == 0)
                
                # 【固定公休特例スラック化 ＆ シワ寄せ防止】
                model.Add(month_sum_var <= target_days)
                model.Add(month_sum_var >= max(0, target_days - cfg.fixed_off_shortage_allowance))
                
                shortage = model.NewIntVar(0, 31, f"shortage_{name}")
                model.Add(shortage == target_days - month_sum_var)
                penalties.append(shortage * WEIGHT_FIXED_OFF_SHORTAGE)
                
            elif s.get('basic_off_type') == 'OR':
                for w in range(6):
                    if w not in s.get('basic_off_days', []):
                        model.Add(y_mo[(name, w)] == 0)
                
                # 指定公休日がある週（ISO週）を特定
                spec_off_iso_weeks = {d.isocalendar()[:2] for d in s.get('specified_off_dates', set()) if d.weekday() != 6 and d not in holidays_jp}
                
                for d in all_dates:
                    if d.weekday() < 6:
                        # 同じ週に指定公休日がある場合、指定公休日がその週の公休となるため、主公休曜日の強制休み義務を免除（振替出勤を許容して総出勤日数を維持）
                        if d.isocalendar()[:2] in spec_off_iso_weeks:
                            continue
                        model.AddImplication(y_mo[(name, d.weekday())], x[(name, d)].Not())

                if is_special:
                    model.Add(month_sum_var >= target_days)
                    model.Add(month_sum_var <= min(target_days + 5, basic_slots - vac_count - sp_vac_count))
                    extra_work = model.NewIntVar(0, 5, f"extra_work_{name}")
                    model.Add(extra_work == month_sum_var - target_days)
                    penalties.append(extra_work * WEIGHT_SPECIAL_EXTRA_WORK)
                else:
                    model.Add(month_sum_var == target_days)

                valid_w = [w for w in range(6) if w in s.get('basic_off_days', [])]
                if valid_w:
                    model.AddExactlyOne([y_mo[(name, w)] for w in valid_w])

        # 連勤上限日数制約 (動的パラメータ)
        k_days = cfg.max_consecutive_work_days
        prev_tail = prev_month_tail_shifts.get(name, [])
        if len(prev_tail) < k_days:
            prev_tail = [0] * (k_days - len(prev_tail)) + prev_tail
        combined_vars = prev_tail + [x[(name, d)] for d in all_dates]
        for i in range(len(combined_vars) - k_days):
            model.Add(sum(combined_vars[i:i + k_days + 1]) <= k_days)

        # 共通ソフトペナルティ
        if s.get('is_core'):
            work_dates = [d for d in all_dates if d.weekday() != 6 and d not in holidays_jp]
            if cfg.weight_off_3consec > 0:
                for i in range(len(work_dates) - 2):
                    d1, d2, d3 = work_dates[i], work_dates[i+1], work_dates[i+2]
                    off3_var = model.NewBoolVar(f"off3_{name}_{d1}")
                    model.Add(x[(name, d1)] + x[(name, d2)] + x[(name, d3)] == 0).OnlyEnforceIf(off3_var)
                    model.Add(x[(name, d1)] + x[(name, d2)] + x[(name, d3)] > 0).OnlyEnforceIf(off3_var.Not())
                    penalties.append(off3_var * cfg.weight_off_3consec)

            if s.get('basic_off_type') == 'OR':
                for d in all_dates:
                    if d.weekday() != 6 and d not in holidays_jp:
                        is_off = model.NewBoolVar(f"is_off_{name}_{d}")
                        model.Add(x[(name, d)] == 0).OnlyEnforceIf(is_off)
                        model.Add(x[(name, d)] == 1).OnlyEnforceIf(is_off.Not())
                        
                        mismatch = model.NewBoolVar(f"mismatch_{name}_{d}")
                        model.AddImplication(is_off, mismatch).OnlyEnforceIf(y_mo[(name, d.weekday())].Not())
                        penalties.append(mismatch * WEIGHT_MAIN_OFF_MISMATCH)

                if name in prev_month_main_offs:
                    prev_w = prev_month_main_offs[name]
                    if prev_w in s.get('basic_off_days', []):
                        penalties.append(y_mo[(name, prev_w)] * WEIGHT_PREV_OFF_REPEAT)

            # 【ローリングウィンドウ化（勤務パターン適応型）】
            if cfg.weekly_off_balance_rule and '週4' in pattern:
                target_off = 3
            else:
                target_off = cfg.weekly_off_target

            prev_tail_7 = prev_month_tail_shifts.get(name, [])
            if len(prev_tail_7) < 6:
                prev_tail_7 = [0] * (6 - len(prev_tail_7)) + prev_tail_7
            hist_vars = prev_tail_7 + [x[(name, d)] for d in all_dates]
            
            for i in range(len(hist_vars) - 6):
                window = hist_vars[i:i+7]
                off_sum = sum(1 - v for v in window)
                diff = model.NewIntVar(-7, 7, f"w_diff_{name}_{i}")
                abs_diff = model.NewIntVar(0, 7, f"w_abs_diff_{name}_{i}")
                model.Add(diff == off_sum - target_off)
                model.AddAbsEquality(abs_diff, diff)
                penalties.append(abs_diff * WEIGHT_ROLLING_WINDOW_DIFF)

            tail_dates = all_dates[-3:]
            tail_off_sum = sum(1 - x[(name, d)] for d in tail_dates)
            penalties.append(tail_off_sum * WEIGHT_TAIL_OFF_SUM)

            # 【カレンダー週単位の公休平準化ペナルティ (v1.6)】
            if cfg.weekly_off_balance_rule:
                for w_key, week_dates in weeks.items():
                    week_work_dates = [d for d in week_dates if d.weekday() != 6 and d not in holidays_jp]
                    if len(week_work_dates) >= 5: # 通常週(6日) または 祝日週(5日)
                        w_off_sum = sum(1 - x[(name, d)] for d in week_work_dates)
                        # 週4日勤務者: 通常週(6日)なら公休目標2日、祝日週(5日)なら公休目標1日
                        # 週5日勤務者: 通常週(6日)なら公休目標1日、祝日週(5日)なら公休目標1日
                        if '週4' in pattern:
                            expected_off = 2 if len(week_work_dates) == 6 else 1
                        else:
                            expected_off = 1
                        
                        w_diff = model.NewIntVar(-6, 6, f"wk_off_diff_{name}_{w_key[0]}_{w_key[1]}")
                        w_abs_diff = model.NewIntVar(0, 6, f"wk_off_abs_{name}_{w_key[0]}_{w_key[1]}")
                        model.Add(w_diff == w_off_sum - expected_off)
                        model.AddAbsEquality(w_abs_diff, w_diff)
                        penalties.append(w_abs_diff * WEIGHT_WEEKLY_OFF_DIFF)

    # 【祝日週の最低1日公休制約 (v1.5)】
    # 祝日により週内勤務可能日数（月〜土の診療日）が5日となる週において、
    # 週5日勤務・週4日勤務を問わず、コアスタッフ全員に最低1日は公休日を設定（出勤日数を最大4日に制限）。
    # これにより、祝日週に5連勤となることを防ぎ、セクション人数平準化ペナルティにより
    # 最も余剰のある平日に自動的に公休（固定公休者の振替公休を含む）が設定される。
    if cfg.holiday_week_off_rule:
        for w_key, week_dates in weeks.items():
            # 週内の診療日 (月〜土 かつ 祝日以外)
            week_work_days = [d for d in week_dates if d.weekday() != 6 and d not in holidays_jp]
            
            # 週内に祝日があり、かつ診療日が5日以上ある週
            has_holiday_in_week = any(d in holidays_jp for d in week_dates)
            if has_holiday_in_week and len(week_work_days) >= 5:
                max_work = len(week_work_days) - 1  # 診療日5日なら最大4日出勤 (最低1日公休)

                for s in staff_list:
                    if not s.get('is_core'):
                        continue
                    if s.get('pattern') == '除外':
                        continue

                    name = s['name']
                    week_work_sum = sum(x[(name, d)] for d in week_work_days)

                    if stage == 1:
                        model.Add(week_work_sum <= max_work)
                    else:
                        slack = model.NewIntVar(0, len(week_work_days), f"sub_slack_{name}_{w_key[0]}_{w_key[1]}")
                        model.Add(week_work_sum - slack <= max_work)
                        penalties.append(slack * WEIGHT_SUBSTITUTE_OFF_SLACK)

    sec_cnt = {}
    sec_core_cnt_vars = {} 
    
    for d in all_dates:
        for sec in ["1階", "2階表", "2階裏"]:
            sec_staffs = [s['name'] for s in staff_list if s['section'] == sec]
            core_staffs = [s['name'] for s in staff_list if s['section'] == sec and s['is_core']]
            
            cnt_var = model.NewIntVar(0, len(sec_staffs), f"sec_cnt_{sec}_{d.strftime('%Y%m%d')}")
            sec_cnt[(sec, d)] = cnt_var
            model.Add(cnt_var == sum(x[(n, d)] for n in sec_staffs)) 
            
            core_cnt_var = model.NewIntVar(0, len(core_staffs) if core_staffs else 1, f"sec_core_cnt_{sec}_{d.strftime('%Y%m%d')}")
            sec_core_cnt_vars[(sec, d)] = core_cnt_var
            if core_staffs:
                model.Add(core_cnt_var == sum(x[(n, d)] for n in core_staffs))
            else:
                model.Add(core_cnt_var == 0)
            
            if d.weekday() != 6 and d not in holidays_jp:
                M_s = ideal_mean[sec]
                
                # 1階はパートも含めた全体人数(cnt_var)、2階はコア要員人数(core_cnt_var)を平準化対象とする
                target_cnt_var = cnt_var if sec == "1階" else core_cnt_var
                
                # 【2階裏特例】動的パラメータ (デフォルト3名)
                if sec == "2階裏":
                    vac_in_sec = sum(1 for s in staff_list if s['section'] == sec and s['is_core'] and (
                        d in s.get('vacation_dates', set()) or d in s.get('special_vacation_dates', set()) or d in s.get('specified_off_dates', set())
                    ))
                    eff_min_cap = max(1, cfg.sec_2f_back_base_cap - vac_in_sec) 
                    max_cap = int(math.ceil(M_s)) + cfg.sec_capacity_tolerance
                elif sec == "1階":
                    # 1階はパートも含めた全スタッフの有休・特休・指定公休を考慮
                    vac_in_sec = sum(1 for s in staff_list if s['section'] == sec and (
                        d in s.get('vacation_dates', set()) or d in s.get('special_vacation_dates', set()) or d in s.get('specified_off_dates', set())
                    ))
                    min_cap = max(1, int(math.floor(M_s)) - cfg.sec_capacity_tolerance)
                    max_cap = int(math.ceil(M_s)) + cfg.sec_capacity_tolerance
                    eff_min_cap = max(1, min_cap - vac_in_sec)
                else:
                    min_cap = max(1, int(math.floor(M_s)) - cfg.sec_capacity_tolerance)
                    max_cap = int(math.ceil(M_s)) + cfg.sec_capacity_tolerance
                    vac_in_sec = sum(1 for s in staff_list if s['section'] == sec and s['is_core'] and (
                        d in s.get('vacation_dates', set()) or d in s.get('special_vacation_dates', set()) or d in s.get('specified_off_dates', set())
                    ))
                    eff_min_cap = max(1, min_cap - vac_in_sec)
                
                if stage == 1:
                    model.Add(target_cnt_var >= eff_min_cap)
                    model.Add(target_cnt_var <= max_cap)
                else:
                    slack_min = model.NewIntVar(0, 20, f"slack_min_sec_{sec}_{d}")
                    slack_max = model.NewIntVar(0, 20, f"slack_max_sec_{sec}_{d}")
                    model.Add(target_cnt_var + slack_min >= eff_min_cap)
                    model.Add(target_cnt_var - slack_max <= max_cap)
                    weight_min = WEIGHT_2F_BACK_MIN_SLACK if sec == "2階裏" else WEIGHT_GENERAL_SEC_SLACK
                    penalties.append(slack_min * weight_min)
                    penalties.append(slack_max * WEIGHT_GENERAL_SEC_SLACK)

                sec_cnt_sq = model.NewIntVar(0, 400, f"sec_cnt_sq_{sec}_{d.strftime('%Y%m%d')}")
                model.AddMultiplicationEquality(sec_cnt_sq, [target_cnt_var, target_cnt_var])
                penalties.append(sec_cnt_sq * WEIGHT_SEC_BALANCE_SQ)
                
                diff = model.NewIntVar(-100, 100, f"diff_{sec}_{d}")
                abs_diff = model.NewIntVar(0, 100, f"abs_diff_{sec}_{d}")
                model.Add(diff == target_cnt_var - int(round(M_s)))
                model.AddAbsEquality(abs_diff, diff)
                penalties.append(abs_diff * WEIGHT_SEC_BALANCE_ABS)

                # 【1階特例】平均目標人数が6.5名以上の場合、できるだけ6名以下を避けるペナルティ
                if sec == "1階" and M_s >= 6.5:
                    is_under_7 = model.NewBoolVar(f"is_under_7_{d.strftime('%Y%m%d')}")
                    model.Add(target_cnt_var <= 6).OnlyEnforceIf(is_under_7)
                    model.Add(target_cnt_var >= 7).OnlyEnforceIf(is_under_7.Not())
                    penalties.append(is_under_7 * 1500000)

        if d.weekday() != 6 and d not in holidays_jp:
            c2b = sec_cnt[("2階裏", d)] 
            is_3 = model.NewBoolVar(f"is_3_{d}")
            model.Add(c2b == 3).OnlyEnforceIf(is_3)
            model.Add(c2b != 3).OnlyEnforceIf(is_3.Not())
            penalties.append(is_3 * WEIGHT_2F_BACK_3_PENALTY)
            
            is_5 = model.NewBoolVar(f"is_5_{d}")
            model.Add(c2b == 5).OnlyEnforceIf(is_5)
            model.Add(c2b != 5).OnlyEnforceIf(is_5.Not())
            penalties.append(is_5 * WEIGHT_2F_BACK_5_PENALTY)
            
        nurse_2f_names = [s['name'] for s in staff_list if s['section'] == '2階裏' and s['job_role'] == '看護師']
        if d.weekday() != 6 and d not in holidays_jp and nurse_2f_names:
            nurse_cnt = sum(x[(n, d)] for n in nurse_2f_names)
            # 最低人数制約 (動的パラメータ)
            if stage == 1:
                model.Add(nurse_cnt >= cfg.nurse_min_count)
            else:
                slack_n = model.NewIntVar(0, len(nurse_2f_names), f"slack_nurse_{d}")
                model.Add(nurse_cnt + slack_n >= cfg.nurse_min_count)
                penalties.append(slack_n * WEIGHT_NURSE_DEFICIT_SLACK)
            
            # 推奨人数制約 (動的パラメータ)
            if cfg.nurse_rec_count > cfg.nurse_min_count:
                is_lt_rec = model.NewBoolVar(f"nurse_lt_rec_{d}")
                model.Add(nurse_cnt < cfg.nurse_rec_count).OnlyEnforceIf(is_lt_rec)
                model.Add(nurse_cnt >= cfg.nurse_rec_count).OnlyEnforceIf(is_lt_rec.Not())
                penalties.append(is_lt_rec * WEIGHT_NURSE_LT2_PENALTY)

    # 【2階裏 通常週の4名体制平準化 (v1.6)】
    # 月〜土の6診療日がある通常週において、特定週に3名体制が集中するのを防止し各週に4名出勤日を均等分散
    if cfg.sec_2f_back_weekly_4_min > 0:
        for w_key, week_dates in weeks.items():
            week_work_days = [d for d in week_dates if d.weekday() != 6 and d not in holidays_jp]
            if len(week_work_days) == 6:
                # 2階裏スタッフの当該週有休・特別有休・指定公休による人数減を考慮
                vac_in_week = sum(
                    1 for d in week_work_days for s in staff_list
                    if s['section'] == '2階裏' and s['is_core'] and (
                        d in s.get('vacation_dates', set()) or d in s.get('special_vacation_dates', set()) or d in s.get('specified_off_dates', set())
                    )
                )
                eff_min_4 = max(0, cfg.sec_2f_back_weekly_4_min - vac_in_week)
                if eff_min_4 > 0:
                    cnt_4_list = []
                    for d in week_work_days:
                        b4 = model.NewBoolVar(f"sec_b4_{w_key[0]}_{w_key[1]}_{d.strftime('%Y%m%d')}")
                        model.Add(sec_cnt[("2階裏", d)] >= 4).OnlyEnforceIf(b4)
                        model.Add(sec_cnt[("2階裏", d)] < 4).OnlyEnforceIf(b4.Not())
                        cnt_4_list.append(b4)
                    
                    if stage == 1:
                        model.Add(sum(cnt_4_list) >= eff_min_4)
                    else:
                        slack_4 = model.NewIntVar(0, 6, f"slack_w4_{w_key[0]}_{w_key[1]}")
                        model.Add(sum(cnt_4_list) + slack_4 >= eff_min_4)
                        penalties.append(slack_4 * WEIGHT_2F_BACK_WEEKLY_4_SLACK)

    for idx, row in df_rules.iterrows():
        r_cond = str(row.get('適用条件', '')).strip()
        r_type = str(row.get('制限タイプ', '')).strip()
        r_staffs_str = str(row.get('対象スタッフ', '')).strip()
        r_min = row.get('最小出勤数')
        r_rec = row.get('推奨出勤数')
        r_lvl = str(row.get('制約レベル', '')).strip().upper()

        if pd.isna(r_staffs_str) or not r_staffs_str: continue
        targets = [t.strip() for t in re.split(r'[、,・]', r_staffs_str) if t.strip()]
        valid_staffs = [n for n in targets if n in [s['name'] for s in staff_list]]
        if not valid_staffs: continue

        try: val_min = int(float(r_min)) if pd.notna(r_min) else 1
        except: val_min = 1
            
        try: val_rec = int(float(r_rec)) if pd.notna(r_rec) else val_min
        except: val_rec = val_min

        target_dates = parse_rule_target_dates(r_cond, all_dates, holidays_jp, ope_dict)

        for d in target_dates:
            if d.weekday() == 6 or d in holidays_jp: continue
            sum_x = sum(x[(n, d)] for n in valid_staffs)
            eff_min = max(val_min, len(valid_staffs) - val_min) if '同時休み' in r_type else val_min
            
            is_hard = 'HARD' in r_lvl
            if is_hard:
                if stage == 1:
                    model.Add(sum_x >= eff_min)
                else:
                    slack_min = model.NewIntVar(0, len(valid_staffs), f"slack_min_{idx}_{d}")
                    model.Add(sum_x + slack_min >= eff_min)
                    weight = LEVEL_WEIGHT_MAP.get(r_lvl, LEVEL_WEIGHT_MAP['HARD'])
                    penalties.append(slack_min * weight)
            else:
                weight = LEVEL_WEIGHT_MAP.get(r_lvl, LEVEL_WEIGHT_MAP['SOFT'])
                slack_min = model.NewIntVar(0, len(valid_staffs), f"slack_min_{idx}_{d}")
                model.Add(sum_x + slack_min >= eff_min)
                penalties.append(slack_min * weight)

            if val_rec > eff_min:
                slack_rec = model.NewIntVar(0, len(valid_staffs), f"slack_rec_{idx}_{d}")
                model.Add(sum_x + slack_rec >= val_rec)
                rec_weight = REC_WEIGHT_MAP.get(r_lvl, 200000)
                penalties.append(slack_rec * rec_weight)

    if penalties:
        model.Minimize(sum(penalties))
        
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = SOLVER_MAX_TIME_SECONDS
    solver.parameters.num_search_workers = SOLVER_NUM_SEARCH_WORKERS
    status = solver.Solve(model)
    
    if status in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        best_x = {}
        for s in staff_list:
            for d in all_dates:
                best_x[(s['name'], d)] = solver.Value(x[(s['name'], d)])
        
        main_off_dict = {}
        for s in staff_list:
            if s.get('is_core'):
                if s.get('basic_off_type') == 'OR':
                    for w in range(6):
                        if solver.Value(y_mo[(s['name'], w)]) == 1:
                            main_off_dict[s['name']] = w
                            break
                elif s.get('basic_off_type') == 'AND':
                    if s.get('basic_off_days'):
                        main_off_dict[s['name']] = s.get('basic_off_days')[0]
            else:
                main_off_dict[s['name']] = -1

        new_tail = {}
        for s in staff_list:
            new_tail[s['name']] = [best_x[(s['name'], d)] for d in all_dates[-7:]] 
            
        return best_x, main_off_dict, target_days_dict, new_tail
    
    return None, None, None, None
