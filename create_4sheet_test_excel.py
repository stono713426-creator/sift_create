"""
第4シート（前月確定シフト実績）付きのテスト用Excelファイルを生成するスクリプト。
榊原さんに9月末尾5連勤（9/26〜9/30）を設定し、10/01の出勤が6連勤防止により
確実に「公休」に制御されるかを検証できるようにする。
"""
import io
import datetime
import pandas as pd
import openpyxl

from create_test_excel import STAFF_MASTER_CSV, RULES_CSV, OPE_CSV

def create_4sheet_excel():
    output_filename = "ope勤務条件表_2026_4sheets.xlsx"
    print(f"Creating {output_filename} with 4 sheets...")

    df_staff = pd.read_csv(io.StringIO(STAFF_MASTER_CSV))
    df_rules = pd.read_csv(io.StringIO(RULES_CSV))
    df_ope = pd.read_csv(io.StringIO(OPE_CSV))

    # 「小西眼科」ルールを除外
    df_rules = df_rules[~df_rules['適用条件'].str.contains('小西')].reset_index(drop=True)

    # 第4シート: 前月シフト実績 (2026年9月分、現場フォーマット縦型)
    wb = openpyxl.Workbook()
    ws_staff = wb.active
    ws_staff.title = "スタッフマスタ"
    
    # 1. スタッフマスタ
    ws_staff.append(df_staff.columns.tolist())
    for r in df_staff.itertuples(index=False):
        ws_staff.append(list(r))

    # 2. シフト制約条件
    ws_rules = wb.create_sheet("シフト制約条件")
    ws_rules.append(df_rules.columns.tolist())
    for r in df_rules.itertuples(index=False):
        ws_rules.append(list(r))

    # 3. 手術スケジュール
    ws_ope = wb.create_sheet("手術スケジュール")
    ws_ope.append(df_ope.columns.tolist())
    for r in df_ope.itertuples(index=False):
        ws_ope.append(list(r))

    # 4. 前月シフト実績 (9月確定シフト)
    ws_prev = wb.create_sheet("前月シフト実績")
    # 行1: タイトル
    ws_prev.append(["2026年9月 勤務シフト表"])
    # 行2: ヘッダー
    staff_names = [s for s in df_staff['氏名'].tolist() if s != '森川']
    headers = ["日付", "曜日", "備考", "警告"] + staff_names
    ws_prev.append(headers)

    # 9月の日付行 (9/01 〜 9/30)
    # 特に 9/26(土), 9/27(日), 9/28(月), 9/29(火), 9/30(水)
    # 榊原さんに 9/26〜9/30 を 5連勤（○）として設定！
    wd_names = ["月", "火", "水", "木", "金", "土", "日"]
    for day in range(1, 31):
        dt = datetime.date(2026, 9, day)
        d_str = dt.strftime("%m/%d")
        wd_str = wd_names[dt.weekday()]
        
        row_vals = [d_str, wd_str, "", ""]
        for s in staff_names:
            if s == "榊原" and day in [26, 27, 28, 29, 30]:
                row_vals.append("○")
            elif dt.weekday() == 6:  # 日曜は休業/公休
                row_vals.append("休業" if s != "榊原" or day not in [26, 27, 28, 29, 30] else "○")
            else:
                row_vals.append("○" if day % 2 == 0 else "公")
        ws_prev.append(row_vals)

    # 集計行（主公休曜日）
    sum_row_label = ["", "", "主公休曜日", ""]
    for s in staff_names:
        sum_row_label.append("月" if s == "小田切" else "土")
    ws_prev.append(sum_row_label)

    wb.save(output_filename)
    wb.close()
    print(f"Successfully generated {output_filename}")

if __name__ == "__main__":
    create_4sheet_excel()
