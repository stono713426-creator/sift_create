import openpyxl
import datetime
import shutil
import sys
from shift_core.generator import ShiftGeneratorService

def test():
    sys.stdout.reconfigure(encoding='utf-8')
    print("=== 指定公休日のHARD適用テスト ===")
    
    # テスト用Excelの作成
    test_file = "test_specified_off_template.xlsx"
    shutil.copy2("ope勤務条件表_2026_4sheets.xlsx", test_file)
    
    wb = openpyxl.load_workbook(test_file)
    ws = wb['スタッフマスタ']
    
    # 石渡さん（1階・正社員）に特別有休(10/14, 10/15) と 指定公休日(10/13, 10/16) を設定
    # 小田切さん（2階表）にも指定公休日(10/20) を設定
    headers = [cell.value for cell in ws[1]]
    name_col = headers.index('氏名') + 1
    sp_vac_col = headers.index('特別有休（指定日）') + 1
    spec_off_col = headers.index('指定公休日') + 1
    
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(row=r, column=name_col).value
        if nm == '石渡':
            ws.cell(row=r, column=sp_vac_col).value = "10/14, 10/15"
            ws.cell(row=r, column=spec_off_col).value = "10/13, 10/16"
            print("石渡: 特別有休=10/14, 10/15, 指定公休日=10/13, 10/16 を設定")
        elif nm == '小田切':
            ws.cell(row=r, column=spec_off_col).value = "10/20"
            print("小田切: 指定公休日=10/20 を設定")
            
    wb.save(test_file)
    wb.close()
    
    # シフト生成
    service = ShiftGeneratorService()
    res = service.generate(input_file=test_file)
    print("生成結果:", res.success)
    assert res.success, "シフト生成に失敗しました"
    
    # 出力結果の確認
    wb_out = openpyxl.load_workbook(res.output_filepath, data_only=False)
    ws_oct = wb_out['2026年10月シフト']
    
    # 列特定
    col_map = {}
    for c in range(5, 40):
        nm = ws_oct.cell(row=2, column=c).value
        if nm: col_map[nm] = c
        
    c_ishiwata = col_map['石渡']
    c_odagiri = col_map['小田切']
    
    # 10/12(月祝), 10/13(火), 10/14(水), 10/15(木), 10/16(金), 10/17(土), 10/18(日)
    print("\n【石渡さんのシフト（10/12〜10/18）】")
    for r in range(3, ws_oct.max_row + 1):
        d_val = ws_oct.cell(row=r, column=1).value
        w_val = ws_oct.cell(row=r, column=2).value
        if not d_val: continue
        
        d_str = str(d_val).split('/')[-1]
        val_ishi = ws_oct.cell(row=r, column=c_ishiwata).value
        val_oda = ws_oct.cell(row=r, column=c_odagiri).value
        
        if d_str in ['12', '13', '14', '15', '16', '17', '18']:
            print(f"10/{d_str}({w_val}): 石渡 = {val_ishi}")
            if d_str == '13':
                assert val_ishi == '公', f"10/13は指定公休日('公')のはずですが {val_ishi} です"
            elif d_str in ['14', '15']:
                assert val_ishi == '特有', f"10/{d_str}は特別有休('特有')のはずですが {val_ishi} です"
            elif d_str == '16':
                assert val_ishi == '公', f"10/16は指定公休日('公')のはずですが {val_ishi} です"
                
        if d_str == '20':
            print(f"\n10/20({w_val}): 小田切 = {val_oda}")
            assert val_oda == '公', f"10/20は小田切の指定公休日('公')のはずですが {val_oda} です"

    print("\n★ すべての指定公休日アサーションに合格しました！")

if __name__ == "__main__":
    test()
